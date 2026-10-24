#!/usr/bin/env python3
"""Check an MD&A variance-commentary workbook against the reference design.

Usage:  python check_mda_workbook.py <workbook.xlsx> [--no-recalc]

Recalculates the workbook with LibreOffice (soffice --headless) so formula results can be read,
then runs the checks below and prints PASS/FAIL per check. Exit code 1 if any check fails.
The reference workbook (SMPL_MDA_Variance_Commentary_Aug2026.xlsx) passes every check.
"""
import os
import re
import shutil
import subprocess
import sys
import tempfile
from collections import Counter

from openpyxl import load_workbook

FRONT = ["Cover", "Executive summary", "NI walk", "Comparison matrix"]
BACK_REQUIRED = ["Watch list", "Tie-outs", "Line data", "Commentary drivers"]
ACTUAL_TABS = ["MoM", "QoQ", "YoY", "YTD"]          # comparisons that all end at the same balance-sheet date
CENTS = re.compile(r"(?<![\d.])\$?\d{1,3}(?:,\d{3})*\.\d{2}(?!\d|%)")
ERR = re.compile(r"^#(REF!|VALUE!|DIV/0!|NAME\?|N/A|NUM!|NULL!)")

results = []


def check(name, ok, detail=""):
    results.append((name, bool(ok), detail))


def recalc(path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    out = tempfile.mkdtemp()
    subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", out, path],
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    p = os.path.join(out, os.path.basename(path))
    return p if os.path.exists(p) else None


def header_row(ws):
    for r in range(1, min(ws.max_row, 12) + 1):
        vals = [str(c.value).strip() if c.value is not None else "" for c in ws[r]]
        if vals and vals[0] == "Line" and "Variance" in vals:
            return r, {v: i + 1 for i, v in enumerate(vals) if v}
    return None, {}


def col_of(heads, prefix):
    for k, v in heads.items():
        if k.startswith(prefix):
            return v
    return None


VIEWS = ["INCOME STATEMENT", "MANAGEMENT P&L", "GL DETAIL", "BALANCE SHEET", "CASH FLOW"]


def band_name(text):
    """A view band row: 'INCOME STATEMENT · accounting view' (reference) or a bare view name in any case."""
    if not isinstance(text, str):
        return None
    head = text.split(" ·")[0].strip()
    if head.upper() in VIEWS:
        return head.upper()
    if len(head) > 3 and head.upper() == head and any(ch.isalpha() for ch in head):
        return head
    return None


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args:
        print(__doc__)
        sys.exit(2)
    path = args[0]
    wb = load_workbook(path)                       # formulas
    rc = None if "--no-recalc" in sys.argv else recalc(path)
    wv = load_workbook(rc, data_only=True) if rc else None   # values
    names = wb.sheetnames

    # 1. Tabs
    check("Tabs: Cover, Executive summary, NI walk, Comparison matrix come first, in that order", names[:4] == FRONT, str(names[:4]))
    missing = [n for n in BACK_REQUIRED if n not in names]
    check("Tabs: Watch list, Tie-outs, Line data and Commentary drivers exist", not missing, f"missing: {missing}" if missing else "")
    end = min([names.index(n) for n in ("Labor bridge", "Watch list") if n in names] or [len(names)])
    comp_tabs = names[4:end]
    check("Tabs: at least one comparison tab", len(comp_tabs) > 0, str(comp_tabs))

    # 2. Formula errors
    if wv:
        errs = [(ws.title, c.coordinate, c.value) for ws in wv.worksheets for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and ERR.match(c.value)]
        check("Formulas: no #REF!, #VALUE!, #DIV/0! or #NAME? after recalculation", not errs, str(errs[:5]))
        # 3. Tie-outs
        cell = None
        if "Tie-outs" in wv.sheetnames:
            for row in wv["Tie-outs"].iter_rows():
                for c in row:
                    if isinstance(c.value, str) and ("All pass" in c.value or "Check the rows" in c.value):
                        cell = c.value
        check("Tie-outs: the summary cell reads \"✓ All pass\"", cell is not None and "All pass" in cell, f"found: {cell!r}")
    else:
        check("Recalculation with LibreOffice", False, "soffice not found or conversion failed; formula checks skipped")

    # 4. Formatting, every sheet (numeric-ness read from recalculated values when available)
    for ws in wb.worksheets:
        vs = wv[ws.title] if wv and ws.title in wv.sheetnames else None
        fonts, general_nums, nums = Counter(), 0, 0
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                fonts[c.font.name] += 1
                v = vs.cell(c.row, c.column).value if vs is not None else c.value
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    if float(v).is_integer() and abs(v) < 10000:
                        continue                      # counts, years, signs
                    nums += 1
                    if c.number_format == "General":
                        general_nums += 1
        top_font, top_n = (fonts.most_common(1) or [("", 0)])[0]
        total = sum(fonts.values()) or 1
        check(f"[{ws.title}] one deliberate font (not the Calibri default)", top_font != "Calibri" and top_n / total > 0.95, f"{dict(fonts.most_common(3))}")
        check(f"[{ws.title}] column widths set", any(d.width for d in ws.column_dimensions.values()))
        if nums:
            check(f"[{ws.title}] numbers carry a number format (not General)", general_nums / nums < 0.05, f"{general_nums} of {nums} numeric cells are General")

    # 5. Comparison tabs
    cur_bs = {}
    for t in comp_tabs:
        ws = wb[t]
        hr, heads = header_row(ws)
        check(f"[{t}] header row with Line … Variance", hr is not None)
        if not hr:
            continue
        cB, cF, cG, cJ, cK = 2, col_of(heads, "F/U"), col_of(heads, "Commentary"), col_of(heads, "Ties"), col_of(heads, "Ledger current")
        cE = col_of(heads, "%")
        lines = sum(1 for r in range(hr + 1, ws.max_row + 1) if ws.cell(r, 2).value is not None)
        check(f"[{t}] has lines (a comparison with no data gets no tab; list it on the Cover instead)", lines > 3, f"{lines} lines")
        check(f"[{t}] freeze panes below the header", ws.freeze_panes is not None)
        check(f"[{t}] ledger-total columns (Ledger current, Ledger comparison) present", cK is not None)
        bands, band = [], None
        rollups = rollups_ok = 0
        fu_on_bs, three_seg, cents, outline = [], [], [], 0
        pct_abs = True
        for r in range(hr + 1, ws.max_row + 1):
            a = ws.cell(r, 1).value
            if band_name(a):
                band = band_name(a)
                bands.append(band)
                continue
            if ws.row_dimensions[r].outlineLevel:
                outline += 1
            if isinstance(a, str) and a.count(" · ") >= 2:
                three_seg.append(f"{r}: {a.strip()[:60]}")
            b = ws.cell(r, cB).value
            if isinstance(b, str) and b.startswith("="):
                rollups += 1
                if cK and ws.cell(r, cK).value is not None:
                    rollups_ok += 1
            if cE:
                e = ws.cell(r, cE).value
                if isinstance(e, str) and e.startswith("=") and "/" in e and "ABS(" not in e.upper():
                    pct_abs = False
            if cF and band and band.startswith("BALANCE"):
                f = ws.cell(r, cF).value
                if f not in (None, ""):
                    fu_on_bs.append(r)
            if cG:
                g = ws.cell(r, cG).value
                if isinstance(g, str) and CENTS.search(g):
                    cents.append(f"{r}: …{CENTS.search(g).group(0)}…")
            if wv and band and band.startswith("BALANCE") and isinstance(a, str) and t in ACTUAL_TABS:
                v = wv[t].cell(r, cB).value
                if isinstance(v, (int, float)):
                    cur_bs.setdefault(a.strip(), {})[t] = round(v, 2)
        check(f"[{t}] a band for each view: {', '.join(VIEWS)}", all(w in bands for w in VIEWS), f"found: {bands}")
        check(f"[{t}] roll-ups are formulas over their lines", rollups > 0, f"{rollups} formula roll-ups")
        check(f"[{t}] every roll-up has the ledger's own total beside it", rollups and rollups_ok == rollups, f"{rollups_ok} of {rollups}")
        check(f"[{t}] GL lines grouped under statement lines (outline levels)", outline > 0, f"{outline} grouped rows")
        check(f"[{t}] % is variance ÷ ABS(comparison)", pct_abs)
        check(f"[{t}] no F/U on balance-sheet lines", not fu_on_bs, f"rows {fu_on_bs[:8]}")
        check(f"[{t}] commentary cites no cents", not cents, "; ".join(cents[:4]))
        check(f"[{t}] line labels have at most two segments", not three_seg, "; ".join(three_seg[:3]))
        if cG:
            check(f"[{t}] commentary column wraps", ws.cell(hr + 3, cG).alignment.wrap_text or any(ws.cell(r, cG).alignment.wrap_text for r in range(hr + 1, min(ws.max_row, hr + 40) + 1)))

    # 6. Balance sheet is point-in-time: the current balance of a line is the same on every actual-vs-actual tab
    if wv and cur_bs:
        bad = {k: v for k, v in cur_bs.items() if len(set(v.values())) > 1}
        check("Balance sheet: current balances agree across MoM, QoQ, YoY and YTD (balances are never summed)", not bad,
              "; ".join(f"{k}: {v}" for k, v in list(bad.items())[:3]))

    # 7. Walks and summary checks (recalculated)
    if wv and "NI walk" in wv.sheetnames:
        ws = wv["NI walk"]
        rows = [r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 1).value or "").startswith("Walk lands")]
        marks = [ws.cell(r, c).value for r in rows for c in range(2, ws.max_column + 1) if ws.cell(r, c).value in ("✓", "✗")]
        check("NI walk: starts at comparison net income, lands on current net income (every column ✓)", rows and marks and all(m == "✓" for m in marks),
              f"{len(rows)} walk(s), marks {Counter(marks)}")
    if wv and "Executive summary" in wv.sheetnames:
        ws = wv["Executive summary"]
        rows = [r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 1).value or "").startswith("Sum of cited")]
        marks = [ws.cell(r, c).value for r in rows for c in range(2, 9) if ws.cell(r, c).value in ("✓", "✗")]
        check("Executive summary: cited net income components add up (every comparison ✓)", rows and marks and all(m == "✓" for m in marks),
              f"{len(rows)} summaries, marks {Counter(marks)}")
    if wv and "Line data" in wv.sheetnames:
        ws = wv["Line data"]
        col = next((c.column for c in ws[1] if c.value == "Drivers tie"), None)
        marks = Counter(ws.cell(r, col).value for r in range(2, ws.max_row + 1)) if col else Counter()
        check("Line data: drivers tie on every line", col and marks.get("✗", 0) == 0 and marks.get("✓", 0) > 0, str(dict(marks)))

    # 8. Cover
    if "Cover" in names:
        ws = wb["Cover"]
        labels = {str(c.value).strip() for row in ws.iter_rows() for c in row if isinstance(c.value, str)}
        need = ["Package ID", "Status", "Data snapshot", "Content hash", "Tie-outs"]
        check("Cover: package stamp (Package ID, Status, Data snapshot, Content hash, Tie-outs)", all(n in labels for n in need),
              f"missing: {[n for n in need if n not in labels]}")
        links = sum(1 for row in ws.iter_rows() for c in row if c.hyperlink)
        check("Cover: contents link to every tab", links >= len(names) - 1, f"{links} links for {len(names) - 1} tabs")

    failed = [r for r in results if not r[1]]
    w = max(len(n) for n, _, _ in results)
    for n, ok, d in results:
        print(f"{'PASS' if ok else 'FAIL'}  {n:<{w}}  {d if not ok else ''}".rstrip())
    print(f"\n{len(results) - len(failed)} of {len(results)} checks pass")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
