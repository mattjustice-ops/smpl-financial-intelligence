"""Render the SMPL MD&A variance-commentary package (mdna_package.json from mdna_export.js) as an Excel workbook.

Ledger amounts are constants (blue); everything derived is a formula (black); links across sheets are green.
Roll-ups recompute from their lines and are checked against the ledger's own totals, so the workbook proves its ties.
"""
import json
import math
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.properties import Outline

PKG = json.load(open(sys.argv[1] if len(sys.argv) > 1 else "mdna_package.json"))
OUT = sys.argv[2] if len(sys.argv) > 2 else "SMPL_MDA_Variance_Commentary_Aug2026.xlsx"

FONT = "Arial"
TEAL, TEAL_LIGHT, SAND, GREY, RULE = "0E5A5F", "E6EFEF", "F2F1EC", "5F6469", "D9D6CE"
BLUE, GREEN, BLACK = "0000FF", "008000", "000000"
AMT = '#,##0.0;(#,##0.0);"–"'
PCT = '0.0%;(0.0%);"–"'
RATIO = '0.00;(0.00);"–"'
thin = Side(style="thin", color=RULE)
VIEW_TITLE = {"IS": "INCOME STATEMENT · accounting view", "MP": "MANAGEMENT P&L · functional view with allocations",
              "GL": "GL DETAIL · natural accounts by department", "BS": "BALANCE SHEET", "CF": "CASH FLOW · indirect"}
K = lambda cents: round(cents / 100000, 5)   # USD thousands  # noqa: E731


def font(color=BLACK, bold=False, size=10, italic=False):
    return Font(name=FONT, color=color, bold=bold, size=size, italic=italic)


def put(ws, r, c, v, color=BLACK, bold=False, size=10, fmt=None, fill=None, wrap=False, align=None, indent=0, italic=False, valign="top"):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font(color, bold, size, italic)
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = PatternFill("solid", start_color=fill)
    cell.alignment = Alignment(wrap_text=wrap, horizontal=align, vertical=valign, indent=indent)
    return cell


def band(ws, r, text, c1=1, c2=12, fill=TEAL, color="FFFFFF", size=10):
    for c in range(c1, c2 + 1):
        ws.cell(row=r, column=c).fill = PatternFill("solid", start_color=fill)
    put(ws, r, c1, text, color=color, bold=True, size=size, fill=fill, valign="center")
    ws.row_dimensions[r].height = 18


def lines_needed(text, width):
    if not text:
        return 1
    per = max(10, int(width * 1.15))
    return sum(max(1, math.ceil(len(part) / per)) for part in str(text).split("\n"))


def ref(sheet, col, row, absolute=False):
    return f"'{sheet}'!{'$' if absolute else ''}{col}{'$' if absolute else ''}{row}"


def ok_mark(b):
    return "✓" if b else "✗"


wb = Workbook()
wb.remove(wb.active)
stamp = PKG["stamp"]
comps = PKG["comparisons"]
ADDR = {}                       # (comp key, view, nid) -> (sheet, row)
CHECKS = {}                     # named check cells for the Tie-outs sheet

ws_cover = wb.create_sheet("Cover")
ws_exec = wb.create_sheet("Executive summary")
ws_walk = wb.create_sheet("NI walk")
ws_matrix = wb.create_sheet("Comparison matrix")
tabs = {c["key"]: wb.create_sheet(c["tab"]) for c in comps}
ws_labor = wb.create_sheet("Labor bridge")
ws_watch = wb.create_sheet("Watch list")
ws_ties = wb.create_sheet("Tie-outs")
ws_data = wb.create_sheet("Line data")
ws_drv = wb.create_sheet("Commentary drivers")

STAMP_LINE = (f"Package {stamp['packageId']} · {stamp['entity']} · {stamp['book']} · {stamp['currency']} thousands · "
              f"Snapshot {stamp['snapshot']} · Content hash {stamp['contentHash']}")


def header(ws, title, sub=None, width_to=12):
    put(ws, 1, 1, title, bold=True, size=14, color=TEAL)
    put(ws, 2, 1, STAMP_LINE, color=GREY, size=9)
    if sub:
        put(ws, 3, 1, sub, color=GREY, size=9)


# ------------------------------------------------------------------------------------------------
# Comparison tabs: every view, statement lines with their GL lines grouped beneath
# ------------------------------------------------------------------------------------------------
COLW = {"A": 46, "B": 12, "C": 12, "D": 12, "E": 8, "F": 5, "G": 82, "H": 56, "I": 15, "J": 6, "K": 12, "L": 12}
for comp in comps:
    ws = tabs[comp["key"]]
    for col, w in COLW.items():
        ws.column_dimensions[col].width = w
    ws.sheet_properties.outlinePr = Outline(summaryBelow=False)
    ap = comp["approval"]
    header(ws, f"{comp['short']} · {comp['label']}",
           f"Approved by {ap['by']} · {ap['at']} · {ap['inScope']} lines in scope, all reviewed · Ledger amounts in blue, formulas in black; "
           f"GL lines are grouped under the statement lines built from them (use the outline buttons to collapse).")
    heads = ["Line", comp["curLabel"], comp["cmpLabel"], "Variance", "%", "F/U", "Commentary (approved)",
             "Future impacts as of Aug 2026 close", "Reviewed", "Ties", "Ledger current", "Ledger comparison"]
    for i, h in enumerate(heads, 1):
        put(ws, 5, i, h, bold=True, color="FFFFFF", fill=TEAL, wrap=True, align="right" if 2 <= i <= 6 or i >= 10 else None, valign="center")
    ws.row_dimensions[5].height = 30
    ws.freeze_panes = "B6"
    r = 6
    for view in ["IS", "MP", "GL", "BS", "CF"]:
        V = comp["views"][view]
        r += 1
        band(ws, r, VIEW_TITLE[view], fill=SAND, color=TEAL)
        if view == "BS" and comp["bsLabels"]:
            put(ws, r, 7, f"Balances: end of {comp['bsLabels'][0]} vs end of {comp['bsLabels'][1]}", color=GREY, size=9, fill=SAND)
        r += 1
        if not V["rows"]:
            put(ws, r, 1, V["na"] + ": no " + ("balance sheet" if view == "BS" else "cash flow statement") + " in the plan version; this comparison runs against prior periods only.",
                color=GREY, italic=True)
            r += 1
            continue
        rows = V["rows"]
        rowmap = {}
        for i, x in enumerate(rows):
            rowmap[x["nid"]] = r + i
        for x in rows:
            rr = rowmap[x["nid"]]
            ADDR[(comp["key"], view, x["nid"])] = (ws.title, rr)
            lvl, strong = x["level"], x["style"] in ("strong", "total")
            put(ws, rr, 1, x["label"], bold=strong, indent=2 if lvl else 0, color=BLACK if not lvl else "2B2F33")
            if x["kind"] == "leaf":
                put(ws, rr, 2, K(x["cur"]), color=BLUE, fmt=AMT, bold=strong)
                put(ws, rr, 3, K(x["cmp"]), color=BLUE, fmt=AMT)
                put(ws, rr, 10, ok_mark(x["driversOk"] and x["paraOk"]), align="center", color=TEAL if x["driversOk"] and x["paraOk"] else "C00000")
            else:
                for col in ("B", "C"):
                    f = "".join(("+" if cf > 0 else "-") + f"{col}{rowmap[k]}" for k, cf in x["children"])
                    put(ws, rr, 2 if col == "B" else 3, "=" + f.lstrip("+"), fmt=AMT, bold=strong)
                put(ws, rr, 11, K(x["cur"]), color=GREY, fmt=AMT, size=9)
                put(ws, rr, 12, K(x["cmp"]), color=GREY, fmt=AMT, size=9)
                para = "TRUE" if (x["paraOk"] and x["driversOk"]) else "FALSE"
                put(ws, rr, 10, f'=IF(AND(ROUND(B{rr}-K{rr},4)=0,ROUND(C{rr}-L{rr},4)=0,{para}),"✓","✗")', align="center", color=TEAL)
            put(ws, rr, 4, f"=B{rr}-C{rr}", fmt=AMT, bold=strong)
            put(ws, rr, 5, f'=IF(C{rr}=0,"n/m",D{rr}/ABS(C{rr}))', fmt=PCT, align="right")
            if x["sign"] == "INCOME":
                put(ws, rr, 6, f'=IF(ROUND(D{rr},5)=0,"",IF(D{rr}>0,"F","U"))', align="center")
            elif x["sign"] == "EXPENSE":
                put(ws, rr, 6, f'=IF(ROUND(D{rr},5)=0,"",IF(D{rr}<0,"F","U"))', align="center")
            text = x["text"] if x["inScope"] or x["level"] == 0 else x["text"]
            put(ws, rr, 7, text, wrap=True, color=BLACK if x["inScope"] else GREY)
            imp = x.get("impact") and x["impact"]["text"]
            if imp:
                put(ws, rr, 8, imp, wrap=True, size=9, color="2B2F33")
            put(ws, rr, 9, f"{x['reviewedBy']} · {x['reviewedAt']}" if x["reviewedBy"] else "Not required", size=9, color=GREY if not x["reviewedBy"] else BLACK)
            if strong:
                for c in range(1, 13):
                    ws.cell(row=rr, column=c).border = Border(top=thin)
            if x["style"] == "strong":
                for c in range(1, 13):
                    ws.cell(row=rr, column=c).fill = PatternFill("solid", start_color="FAF9F6")
            if lvl:
                ws.row_dimensions[rr].outline_level = 1
            ng = math.ceil(len(text) / (COLW["G"] * 1.22)) if text else 1
            nh = math.ceil(len(imp) / (COLW["H"] * 1.38)) if imp else 1
            ws.row_dimensions[rr].height = max(15, 12.6 * ng + 3, 11.4 * nh + 3)
        r += len(rows)
    r += 1
    put(ws, r, 1, "F = favorable, U = unfavorable. Variance = current − comparison; % is of the comparison amount. Ties: roll-ups recompute from their lines and equal the "
                  "ledger total (columns K–L); GL lines tie their drivers and paragraph at export (T22, T26). Paragraph amounts are rounded so they add to the displayed variance.",
        color=GREY, size=9, wrap=False)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "5:5"
    CHECKS[("tab", comp["key"])] = f"IF(COUNTIF('{ws.title}'!J:J,\"✗\")=0,\"✓\",\"✗\")"

# ------------------------------------------------------------------------------------------------
# Line data and commentary drivers (flat tables), with driver sums recomputed by formula
# ------------------------------------------------------------------------------------------------
dh = ["Key", "Comparison", "Comparison label", "View", "Line ID", "Line", "Level", "Current", "Comparison value", "Variance (ledger)",
      "Variance shown", "Commentary", "Source", "In scope", "Reviewed by", "Reviewed at", "Σ drivers exact", "Σ drivers shown", "Drivers tie"]
for i, h in enumerate(dh, 1):
    put(ws_data, 1, i, h, bold=True, color="FFFFFF", fill=TEAL, wrap=True, valign="center")
drv_h = ["Key", "Driver #", "Driver or component line", "Amount exact", "Amount shown"]
for i, h in enumerate(drv_h, 1):
    put(ws_drv, 1, i, h, bold=True, color="FFFFFF", fill=TEAL, valign="center")
dr, vr = 2, 2
for comp in comps:
    for view in ["IS", "MP", "GL", "BS", "CF"]:
        rows = comp["views"][view]["rows"]
        if not rows:
            continue
        for x in rows:
            key = f"{comp['cid']}|{view}|{x['nid']}"
            vals = [key, comp["short"], comp["label"], view, x["nid"], x["label"], x["level"], K(x["cur"]), K(x["cmp"]), K(x["var"]),
                    round(x["varShown"] / 100000, 1), x["text"], x["source"],
                    "Yes" if x["inScope"] else "No", x["reviewedBy"] or "", x["reviewedAt"] or ""]
            for i, v in enumerate(vals, 1):
                put(ws_data, dr, i, v, color=BLUE if i in (8, 9, 10, 11) else BLACK, fmt=AMT if i in (8, 9, 10, 11) else None)
            put(ws_data, dr, 17, f"=SUMIFS('Commentary drivers'!$D:$D,'Commentary drivers'!$A:$A,A{dr})", fmt=AMT)
            put(ws_data, dr, 18, f"=SUMIFS('Commentary drivers'!$E:$E,'Commentary drivers'!$A:$A,A{dr})", fmt=AMT)
            put(ws_data, dr, 19, f'=IF(AND(ROUND(Q{dr}-J{dr},4)=0,ROUND(R{dr}-K{dr},1)=0),"✓","✗")', align="center")
            for j, d in enumerate(x["drivers"], 1):
                put(ws_drv, vr, 1, key)
                put(ws_drv, vr, 2, j)
                put(ws_drv, vr, 3, d["label"])
                put(ws_drv, vr, 4, K(d["exact"]), color=BLUE, fmt="#,##0.00000;(#,##0.00000);\"–\"")
                put(ws_drv, vr, 5, round(d["shown"] / 100000, 1), color=BLUE, fmt=AMT)
                vr += 1
            dr += 1
for col, w in zip("ABCDEFGHIJKLMNOPQRS", [34, 16, 30, 6, 30, 40, 6, 11, 11, 11, 11, 90, 16, 8, 10, 11, 11, 11, 8]):
    ws_data.column_dimensions[col].width = w
for col, w in zip("ABCDE", [40, 8, 50, 14, 12]):
    ws_drv.column_dimensions[col].width = w
ws_data.freeze_panes = "B2"
ws_drv.freeze_panes = "B2"
ws_data.auto_filter.ref = f"A1:S{dr - 1}"
ws_drv.auto_filter.ref = f"A1:E{vr - 1}"
DATA_LAST = dr - 1
put(ws_drv, vr + 1, 1, "Each line's commentary cites these rows: the top three drivers (or component lines for a statement subtotal) and Other. "
                       "Exact amounts sum to the line's variance in cents (T22); shown amounts sum to the displayed variance (T26). Line data column S checks both.",
    color=GREY, size=9)

# ------------------------------------------------------------------------------------------------
# Executive summary
# ------------------------------------------------------------------------------------------------
ws = ws_exec
header(ws, "Executive summary · August 2026",
       "Drafted by AI from approved lines only. Every amount is bound to a line in this package, and the net income components it cites add up to the "
       "net income change (checked below each summary).")
for col, w in zip("ABCDEFGH", [52, 14, 14, 14, 10, 8, 14, 30]):
    ws.column_dimensions[col].width = w
r = 5
KPI = [("IS", "IS.REV.TOTAL"), ("IS", "IS.GP"), (None, "GM"), ("IS", "IS.OPX.TOTAL"), ("IS", "IS.OI"), ("IS", "IS.NI"), ("MP", "MP.CM2")]
for comp in comps:
    sheet = comp["tab"]
    band(ws, r, f"{comp['short']} · {comp['label']}", c2=8)
    put(ws, r, 6, f"Approved · {comp['approval']['by']} · {comp['approval']['at']}", color="FFFFFF", size=9, fill=TEAL, valign="center")
    r += 1
    s = comp["summary"]
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    put(ws, r, 1, s["text"], wrap=True, size=10.5)
    ws.row_dimensions[r].height = 14 * math.ceil(len(s["text"]) / (166 * 1.15)) + 6
    r += 1
    put(ws, r, 1, "AI draft · validated: no number outside a bound token; net income components cited once each.", color=GREY, size=8.5, italic=True)
    r += 2
    for i, h in enumerate(["Key lines", comp["curLabel"], comp["cmpLabel"], "Variance", "%", "F/U"], 1):
        put(ws, r, i, h, bold=True, align="right" if i > 1 else None, fill=TEAL_LIGHT)
    r += 1
    rev_row = gp_row = None
    for v, nid in KPI:
        if nid == "GM":
            put(ws, r, 1, "Gross margin %", indent=1)
            put(ws, r, 2, f"=IF(B{rev_row}=0,0,B{gp_row}/B{rev_row})", fmt=PCT)
            put(ws, r, 3, f"=IF(C{rev_row}=0,0,C{gp_row}/C{rev_row})", fmt=PCT)
            put(ws, r, 4, f"=B{r}-C{r}", fmt='+0.0%;-0.0%;"–"')
            put(ws, r, 5, "pts", color=GREY, align="right", size=9)
            r += 1
            continue
        sh, row = ADDR[(comp["key"], v, nid)]
        lbl = {"IS.REV.TOTAL": "Total net revenue", "IS.GP": "Gross profit", "IS.OPX.TOTAL": "Total operating expenses", "IS.OI": "Operating income",
               "IS.NI": "Net income", "MP.CM2": "Contribution after direct go-to-market (management view)"}[nid]
        put(ws, r, 1, lbl, bold=nid == "IS.NI")
        put(ws, r, 2, f"={ref(sh, 'B', row)}", color=GREEN, fmt=AMT, bold=nid == "IS.NI")
        put(ws, r, 3, f"={ref(sh, 'C', row)}", color=GREEN, fmt=AMT)
        put(ws, r, 4, f"=B{r}-C{r}", fmt=AMT, bold=nid == "IS.NI")
        put(ws, r, 5, f'=IF(C{r}=0,"n/m",D{r}/ABS(C{r}))', fmt=PCT, align="right")
        put(ws, r, 6, f"={ref(sh, 'F', row)}", color=GREEN, align="center")
        if nid == "IS.REV.TOTAL":
            rev_row = r
        if nid == "IS.GP":
            gp_row = r
        if nid == "IS.NI":
            ni_row = r
        r += 1
    r += 1
    put(ws, r, 1, "Net income components cited in the summary", bold=True, fill=TEAL_LIGHT)
    put(ws, r, 2, "Shown", bold=True, align="right", fill=TEAL_LIGHT)
    r += 1
    first = r
    for d in s["cited"]:
        put(ws, r, 1, d["label"], indent=1)
        put(ws, r, 2, round(d["shown"] / 100000, 1), color=BLUE, fmt=AMT)
        r += 1
    put(ws, r, 1, "Sum of cited components", bold=True)
    put(ws, r, 2, f"=SUM(B{first}:B{r - 1})", fmt=AMT, bold=True)
    put(ws, r, 3, "Net income change", color=GREY, align="right", size=9)
    put(ws, r, 4, f"=ROUND(D{ni_row},1)", fmt=AMT)
    put(ws, r, 5, f'=IF(ROUND(B{r}-D{r},1)=0,"✓","✗")', align="center", bold=True, color=TEAL)
    CHECKS[("exec", comp["key"])] = f"'Executive summary'!E{r}"
    r += 3
ws.page_setup.orientation = "portrait"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

# ------------------------------------------------------------------------------------------------
# NI walk (income statement and management view)
# ------------------------------------------------------------------------------------------------
ws = ws_walk
header(ws, "Net income walk · comparison net income to August 2026",
       "Each step is a statement subtotal's variance times its sign in net income (the same component lines the net income commentary cites). "
       "Steps link to the comparison tabs; the walk must land on the statement's net income.")
ws.column_dimensions["A"].width = 44
for j in range(len(comps)):
    ws.column_dimensions[L(2 + j)].width = 14
IS_STEPS = [("IS.REV.TOTAL", 1, "Total net revenue"), ("IS.COR.TOTAL", -1, "Total cost of revenue"), ("IS.OPX.TOTAL", -1, "Total operating expenses"),
            ("IS.NOP.INT_INC", 1, "Interest income"), ("IS.NOP.INT_EXP", -1, "Interest expense"), ("IS.TAX", -1, "Income tax")]
MP_STEPS = [("MP.REV", 1, "Total net revenue"), ("MP.COR", -1, "Service delivery cost"), ("MP.SM_DIRECT", -1, "S&M — direct go-to-market"),
            ("MP.OPX", -1, "Shared S&M, R&D and G&A (with allocations)"), ("MP.NOP.INT_INC", 1, "Interest income"), ("MP.NOP.INT_EXP", -1, "Interest expense"),
            ("MP.TAX", -1, "Income tax")]
WALK_ROWS = {}
r = 5
for view, steps, title, root in (("IS", IS_STEPS, "Income statement walk", "IS.NI"), ("MP", MP_STEPS, "Management P&L walk (allocations included)", "MP.NI")):
    band(ws, r, title, c2=1 + len(comps))
    r += 1
    put(ws, r, 1, "USD thousands", color=GREY, size=9, fill=TEAL_LIGHT)
    for j, comp in enumerate(comps):
        put(ws, r, 2 + j, comp["short"], bold=True, align="right", fill=TEAL_LIGHT, wrap=True)
    ws.row_dimensions[r].height = 28
    r += 1
    start_row = r
    put(ws, r, 1, "Comparison net income", bold=True)
    for j, comp in enumerate(comps):
        sh, row = ADDR[(comp["key"], view, root)]
        put(ws, r, 2 + j, f"={ref(sh, 'C', row)}", color=GREEN, fmt=AMT, bold=True)
    r += 1
    for nid, sg, lbl in steps:
        put(ws, r, 1, ("+ " if sg > 0 else "− ") + lbl, indent=1)
        for j, comp in enumerate(comps):
            sh, row = ADDR[(comp["key"], view, nid)]
            put(ws, r, 2 + j, f"={'-' if sg < 0 else ''}{ref(sh, 'D', row)}", color=GREEN, fmt=AMT)
        r += 1
    end_row = r
    put(ws, r, 1, "Current net income (walk)", bold=True)
    for j in range(len(comps)):
        c = L(2 + j)
        put(ws, r, 2 + j, f"={c}{start_row}+SUM({c}{start_row + 1}:{c}{r - 1})", fmt=AMT, bold=True)
        ws.cell(row=r, column=2 + j).border = Border(top=thin)
    r += 1
    put(ws, r, 1, "Current net income (statement)", color=GREY)
    for j, comp in enumerate(comps):
        sh, row = ADDR[(comp["key"], view, root)]
        put(ws, r, 2 + j, f"={ref(sh, 'B', row)}", color=GREEN, fmt=AMT)
    r += 1
    put(ws, r, 1, "Walk lands on net income", color=GREY)
    for j, comp in enumerate(comps):
        c = L(2 + j)
        put(ws, r, 2 + j, f'=IF(ROUND({c}{r - 2}-{c}{r - 1},4)=0,"✓","✗")', align="right", color=TEAL, bold=True)
        CHECKS[("walk", view, comp["key"])] = f"'NI walk'!{c}{r}"
    WALK_ROWS[view] = (start_row, end_row, steps)
    r += 2

ws.freeze_panes = "B5"

# ------------------------------------------------------------------------------------------------
# Comparison matrix: variance of every statement line across all comparisons (links)
# ------------------------------------------------------------------------------------------------
ws = ws_matrix
header(ws, "Comparison matrix · variance by line across every comparison",
       "Links to the comparison tabs. Statement lines only; open a tab for GL lines and commentary.")
ws.column_dimensions["A"].width = 46
put(ws, 5, 1, "Line (USD thousands)", bold=True, color="FFFFFF", fill=TEAL)
for j, comp in enumerate(comps):
    c1 = 2 + 2 * j
    ws.merge_cells(start_row=5, start_column=c1, end_row=5, end_column=c1 + 1)
    put(ws, 5, c1, comp["short"], bold=True, color="FFFFFF", fill=TEAL, align="center")
    ws.cell(row=5, column=c1 + 1).fill = PatternFill("solid", start_color=TEAL)
    put(ws, 6, c1, "Variance", bold=True, align="right", fill=TEAL_LIGHT, size=9)
    put(ws, 6, c1 + 1, "%", bold=True, align="right", fill=TEAL_LIGHT, size=9)
    ws.column_dimensions[L(c1)].width = 11
    ws.column_dimensions[L(c1 + 1)].width = 8
put(ws, 6, 1, "", fill=TEAL_LIGHT)
r = 7
for view in ("IS", "MP"):
    band(ws, r, VIEW_TITLE[view], c2=1 + 2 * len(comps), fill=SAND, color=TEAL)
    r += 1
    for x in comps[0]["views"][view]["rows"]:
        if x["level"]:
            continue
        strong = x["style"] in ("strong", "total")
        put(ws, r, 1, x["label"], bold=strong)
        for j, comp in enumerate(comps):
            sh, row = ADDR[(comp["key"], view, x["nid"])]
            put(ws, r, 2 + 2 * j, f"={ref(sh, 'D', row)}", color=GREEN, fmt=AMT, bold=strong)
            put(ws, r, 3 + 2 * j, f"={ref(sh, 'E', row)}", color=GREEN, fmt=PCT, align="right")
        r += 1
    r += 1
ws.freeze_panes = "B7"

# ------------------------------------------------------------------------------------------------
# Labor bridge (§18)
# ------------------------------------------------------------------------------------------------
ws = ws_labor
header(ws, "Compensation bridge (§18) · month over month and vs budget",
       "Bars are the labor engine's (blue). Start and end link to the GL compensation lines; the bridge must land on the ledger. Hires, departures and mix are computed "
       "per department; only their sum (headcount & mix) adds across departments.")
ws.column_dimensions["A"].width = 40
for i in range(2, 17):
    ws.column_dimensions[L(i)].width = 12
r = 5
BARS = [("HIRES", "Hires"), ("DEPARTURES", "Departures"), ("OPENING_HC", "Opening heads vs plan"), ("MIX", "Mix"), ("TRANSFERS_NET", "Transfers"),
        ("RATE", "Rate"), ("TIMING", "Timing"), ("BENEFITS_TAXES", "Benefits & taxes"), ("ONE_TIME", "One-time"), ("GL_RECON", "GL vs payroll")]
for lb, ckey in zip(PKG["labor"], ("MOM", "BUD_M")):
    sheet = next(c for c in comps if c["key"] == ckey)["tab"]
    band(ws, r, lb["title"], c2=16)
    r += 1
    put(ws, r, 1, "Company (all departments)", bold=True, fill=TEAL_LIGHT)
    put(ws, r, 2, "USD thousands", color=GREY, size=9, fill=TEAL_LIGHT, align="right")
    r += 1
    _, comp_row = ADDR[(ckey, "GL", "GL.COMP")]
    s_row = r
    put(ws, r, 1, "Start: " + ("Jul 2026 actual" if ckey == "MOM" else "Budget 2026 v1, Aug") + " (GL 6000)")
    put(ws, r, 2, f"={ref(sheet, 'C', comp_row)}", color=GREEN, fmt=AMT)
    r += 1
    for b in lb["company"]["bars"]:
        put(ws, r, 1, b["label"], indent=1)
        put(ws, r, 2, K(b["amt"]), color=BLUE, fmt=AMT)
        r += 1
    put(ws, r, 1, "End: Aug 2026 actual (bridge)", bold=True)
    put(ws, r, 2, f"=B{s_row}+SUM(B{s_row + 1}:B{r - 1})", fmt=AMT, bold=True)
    ws.cell(row=r, column=2).border = Border(top=thin)
    e_row = r
    r += 1
    put(ws, r, 1, "Aug 2026 compensation (ledger)", color=GREY)
    put(ws, r, 2, f"={ref(sheet, 'B', comp_row)}", color=GREEN, fmt=AMT)
    put(ws, r, 3, f'=IF(ROUND(B{e_row}-B{r},4)=0,"✓","✗")', color=TEAL, bold=True, align="center")
    CHECKS[("labor", ckey)] = f"'Labor bridge'!C{r}"
    hm_company_row = s_row + 1 + [b["key"] for b in lb["company"]["bars"]].index("HEADCOUNT_MIX") if any(b["key"] == "HEADCOUNT_MIX" for b in lb["company"]["bars"]) else None
    r += 2
    put(ws, r, 1, "By department", bold=True, fill=TEAL_LIGHT)
    heads = ["Start"] + [b[1] for b in BARS] + ["End (bridge)", "End (ledger)", "Ties"]
    for i, h in enumerate(heads, 2):
        put(ws, r, i, h, bold=True, align="right", fill=TEAL_LIGHT, wrap=True, size=9)
    ws.row_dimensions[r].height = 28
    r += 1
    d_first = r
    for d in lb["byDept"]:
        _, drow = ADDR[(ckey, "GL", "GL.6000." + d["dept"])]
        put(ws, r, 1, d["name"], indent=1)
        put(ws, r, 2, f"={ref(sheet, 'C', drow)}", color=GREEN, fmt=AMT)
        bars = {b["key"]: b["amt"] for b in d["bars"]}
        for i, (bk, _) in enumerate(BARS, 3):
            put(ws, r, i, K(bars.get(bk, 0)), color=BLUE, fmt=AMT)
        put(ws, r, 13, f"=B{r}+SUM(C{r}:L{r})", fmt=AMT, bold=True)
        put(ws, r, 14, f"={ref(sheet, 'B', drow)}", color=GREEN, fmt=AMT)
        put(ws, r, 15, f'=IF(ROUND(M{r}-N{r},4)=0,"✓","✗")', color=TEAL, bold=True, align="center")
        r += 1
    put(ws, r, 1, "All departments", bold=True)
    for i in range(2, 15):
        put(ws, r, i, f"=SUM({L(i)}{d_first}:{L(i)}{r - 1})", fmt=AMT, bold=True)
        ws.cell(row=r, column=i).border = Border(top=thin)
    put(ws, r, 15, f'=IF(ROUND(M{r}-B{e_row},4)=0,"✓","✗")', color=TEAL, bold=True, align="center")
    r += 1
    if hm_company_row:
        put(ws, r, 1, "Headcount & mix: Σ departments (hires + departures + opening + mix + transfers) = company", color=GREY, size=9)
        put(ws, r, 15, f'=IF(ROUND(C{r - 1}+D{r - 1}+E{r - 1}+F{r - 1}+G{r - 1}-B{hm_company_row},4)=0,"✓","✗")', color=TEAL, bold=True, align="center")
        CHECKS[("labor_hm", ckey)] = f"'Labor bridge'!O{r}"
        r += 1
    r += 2

# ------------------------------------------------------------------------------------------------
# Watch list: next month and beyond, then how last month's watch list played out
# ------------------------------------------------------------------------------------------------
ws = ws_watch
W = PKG["watch"]
header(ws, f"What to watch · future impacts noted at the {W['asOf']} close",
       "Expected change on each GL line versus August, in the line's reported sign (revenue and costs positive). The net income row applies each line's sign. "
       "Suggested items are proposals awaiting acceptance; accepted items flow to the forecast as proposed adjustments.")
widths = [44, 34, 15, 18, 10, 11, 30, 7]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[L(i)].width = w
months = W["months"]
for j in range(len(months)):
    ws.column_dimensions[L(9 + j)].width = 10
r = 5
heads = ["Impact", "GL line", "Type", "Owner", "Confidence", "Status", "Source", "NI sign"] + [m["label"].replace(" 20", " ’") for m in months]
for i, h in enumerate(heads, 1):
    put(ws, r, i, h, bold=True, color="FFFFFF", fill=TEAL, align="right" if i >= 8 else None, wrap=True, valign="center")
ws.row_dimensions[r].height = 28
r += 1
w_first = r
for it in W["items"]:
    vals = [it["title"], it["line"], it["type"], it["owner"], it["confidence"], it["status"], it["source"]]
    for i, v in enumerate(vals, 1):
        put(ws, r, i, v, size=9.5, wrap=i in (1, 7), color="6E4700" if (i == 6 and v == "Suggested") else BLACK)
    put(ws, r, 8, it["niSign"], color=BLUE, align="right")
    for j, m in enumerate(months):
        if m["key"] in it["amounts"]:
            put(ws, r, 9 + j, K(it["amounts"][m["key"]]), color=BLUE, fmt=AMT)
    ws.row_dimensions[r].height = max(15, 12.5 * max(lines_needed(it["title"], 44), lines_needed(it["source"], 30)) + 2)
    r += 1
w_last = r - 1
put(ws, r, 1, "Net income effect by month (formula)", bold=True)
for j in range(len(months)):
    c = L(9 + j)
    put(ws, r, 9 + j, f"=SUMPRODUCT($H${w_first}:$H${w_last},{c}{w_first}:{c}{w_last})", fmt=AMT, bold=True)
    ws.cell(row=r, column=9 + j).border = Border(top=thin)
ni_row = r
r += 1
put(ws, r, 1, "Net income effect (SMPL: GL, income statement and management P&L agree)", color=GREY)
for j, m in enumerate(months):
    put(ws, r, 9 + j, K(W["niByMonth"].get(m["key"], 0)), color=GREY, fmt=AMT)
r += 1
put(ws, r, 1, "Ties", color=GREY)
for j in range(len(months)):
    c = L(9 + j)
    put(ws, r, 9 + j, f'=IF(ROUND({c}{ni_row}-{c}{r - 1},4)=0,"✓","✗")', color=TEAL, bold=True, align="right")
CHECKS[("watch",)] = f"IF(COUNTIF('Watch list'!I{r}:{L(8 + len(months))}{r},\"✗\")=0,\"✓\",\"✗\")"
r += 1
put(ws, r, 1, f"Next month ({months[0]['label']}): net income effect", bold=True, fill=TEAL_LIGHT)
put(ws, r, 9, f"=I{ni_row}", fmt=AMT, bold=True, fill=TEAL_LIGHT)
r += 3

S = PKG["scoring"]
band(ws, r, f"How the {S['notedAt']}-close watch list played out in {S['month']}", c2=14)
r += 1
put(ws, r, 1, "Realized when actual ÷ expected is at or above", color=BLACK)
put(ws, r, 2, 0.9, color=BLUE, fmt="0%", fill="FFFF00")
thr1 = f"$B${r}"
r += 1
put(ws, r, 1, "Partly realized when at or above (else missed)")
put(ws, r, 2, 0.25, color=BLUE, fmt="0%", fill="FFFF00")
thr2 = f"$B${r}"
put(ws, r, 3, "Policy PD-32 (illustrative defaults); edit the yellow cells to test other thresholds.", color=GREY, size=9)
r += 2
sh = ["Impact noted at close", "GL line", "Owner", "Expected", "Actual", "Actual ÷ expected", "Outcome", "SMPL outcome", "Driver matched", "", "", "Follow-up for commentary"]
for i, h in enumerate(sh, 1):
    if h:
        put(ws, r, i, h, bold=True, color="FFFFFF", fill=TEAL, wrap=True, valign="center", align="right" if 4 <= i <= 6 else None)
    else:
        ws.cell(row=r, column=i).fill = PatternFill("solid", start_color=TEAL)
ws.row_dimensions[r].height = 28
r += 1
s_first = r
for x in S["items"]:
    put(ws, r, 1, x["title"], wrap=True, size=9.5)
    put(ws, r, 2, x["lineLabel"], size=9.5, wrap=True)
    put(ws, r, 3, x["owner"], size=9.5)
    put(ws, r, 4, K(x["expected"]), color=BLUE, fmt=AMT)
    put(ws, r, 5, K(x["actual"]), color=BLUE, fmt=AMT)
    put(ws, r, 6, f"=IF(D{r}=0,0,E{r}/D{r})", fmt=RATIO)
    put(ws, r, 7, f'=IF(F{r}>={thr1},"Realized",IF(F{r}>={thr2},"Partly realized","Missed"))', bold=True)
    put(ws, r, 8, x["outcome"], color=GREY, size=9)
    put(ws, r, 9, x["driver"] or "No driver moved", size=9.5)
    if x["outcome"] != "Realized":
        q = "“" + x["title"] + "”"
        fu = (f"Noted at {S['notedAt'][:3]} close: {q} was only partly realized." if x["outcome"] == "Partly realized"
              else f"Noted at {S['notedAt'][:3]} close: {q} did not happen." if x["actual"] == 0
              else f"Noted at {S['notedAt'][:3]} close: {q} did not happen; the line moved the other way.")
        put(ws, r, 12, fu, size=9.5, wrap=True)
    ws.row_dimensions[r].height = 26
    r += 1
s_last = r - 1
put(ws, r, 1, "Outcomes match SMPL", color=GREY)
put(ws, r, 7, f'=IF(SUMPRODUCT(--(G{s_first}:G{s_last}<>H{s_first}:H{s_last}))=0,"✓","✗")', color=TEAL, bold=True)
CHECKS[("score",)] = f"'Watch list'!G{r}"
r += 2
put(ws, r, 1, "Realized", indent=1)
put(ws, r, 2, f'=COUNTIF(G{s_first}:G{s_last},"Realized")')
r += 1
put(ws, r, 1, "Partly realized", indent=1)
put(ws, r, 2, f'=COUNTIF(G{s_first}:G{s_last},"Partly realized")')
r += 1
put(ws, r, 1, "Missed", indent=1)
put(ws, r, 2, f'=COUNTIF(G{s_first}:G{s_last},"Missed")')
r += 1
put(ws, r, 1, "Dollar accuracy: 1 − Σ|actual − expected| ÷ Σ|expected|", bold=True)
put(ws, r, 2, f"=1-SUMPRODUCT(ABS(E{s_first}:E{s_last}-D{s_first}:D{s_last}))/SUMPRODUCT(ABS(D{s_first}:D{s_last}))", fmt="0.0%", bold=True)
r += 2
for i, h in enumerate(["Accuracy by owner", "Items", "Realized", "Partly", "Missed", "Accuracy"], 1):
    put(ws, r, i, h, bold=True, fill=TEAL_LIGHT, align="right" if i > 1 else None)
r += 1
for owner in sorted(S["byOwner"]):
    put(ws, r, 1, owner, indent=1)
    put(ws, r, 2, f'=COUNTIF($C${s_first}:$C${s_last},A{r})')
    put(ws, r, 3, f'=COUNTIFS($C${s_first}:$C${s_last},A{r},$G${s_first}:$G${s_last},"Realized")')
    put(ws, r, 4, f'=COUNTIFS($C${s_first}:$C${s_last},A{r},$G${s_first}:$G${s_last},"Partly realized")')
    put(ws, r, 5, f'=COUNTIFS($C${s_first}:$C${s_last},A{r},$G${s_first}:$G${s_last},"Missed")')
    put(ws, r, 6, f"=1-SUMPRODUCT(($C${s_first}:$C${s_last}=A{r})*ABS($E${s_first}:$E${s_last}-$D${s_first}:$D${s_last}))"
                  f"/SUMPRODUCT(($C${s_first}:$C${s_last}=A{r})*ABS($D${s_first}:$D${s_last}))", fmt="0.0%")
    r += 1
ws.freeze_panes = "B6"

# ------------------------------------------------------------------------------------------------
# Tie-outs: SMPL's validations and the workbook's own recomputation, per comparison
# ------------------------------------------------------------------------------------------------
ws = ws_ties
header(ws, "Tie-outs · every comparison", "SMPL validations at export (T22–T27) and the same ties recomputed by this workbook's formulas.")
ws.column_dimensions["A"].width = 70
put(ws, 5, 1, "Check", bold=True, color="FFFFFF", fill=TEAL)
for j, comp in enumerate(comps):
    put(ws, 5, 2 + j, comp["short"], bold=True, color="FFFFFF", fill=TEAL, align="center", wrap=True)
    ws.column_dimensions[L(2 + j)].width = 13
ws.row_dimensions[5].height = 30
r = 6
band(ws, r, "SMPL validations at export", c2=1 + len(comps), fill=SAND, color=TEAL)
r += 1
ENGINE_LABELS = ["Drivers sum to every line's variance (T22)", "Roll-ups equal the sum of their lines (T23)", "Commentary amounts add up to each variance (T26)",
                 "Net income agrees: GL = income statement = management P&L (T24)", "Allocations net to zero; operating income agrees (T24)",
                 "Balance sheet balances (T27)", "Cash flow ties to the change in cash (T27)"]
t_first = r
for i, lbl in enumerate(ENGINE_LABELS):
    put(ws, r, 1, lbl)
    for j, comp in enumerate(comps):
        ck = comp["checks"][i]
        put(ws, r, 2 + j, "n/a" if ck.get("na") else ok_mark(ck["ok"]), align="center", color=GREY if ck.get("na") else TEAL, bold=not ck.get("na"))
    r += 1
band(ws, r, "Recomputed in this workbook", c2=1 + len(comps), fill=SAND, color=TEAL)
r += 1
WB_ROWS = [
    ("Comparison tab: every roll-up recomputes from its lines to the ledger total; every line ties", lambda c: "=" + CHECKS[("tab", c["key"])]),
    ("Line data: exact driver sums equal each line's variance; shown amounts equal the displayed variance",
     lambda c: f"=IF(COUNTIFS('Line data'!$B$2:$B${DATA_LAST},\"{c['short']}\",'Line data'!$S$2:$S${DATA_LAST},\"✗\")=0,\"✓\",\"✗\")"),
    ("Executive summary: cited net income components add up to the net income change", lambda c: "=" + CHECKS[("exec", c["key"])]),
    ("Net income walk lands on the statement (income statement)", lambda c: "=" + CHECKS[("walk", "IS", c["key"])]),
    ("Net income walk lands on the statement (management P&L, allocations included)", lambda c: "=" + CHECKS[("walk", "MP", c["key"])]),
]
for lbl, fn in WB_ROWS:
    put(ws, r, 1, lbl)
    for j, comp in enumerate(comps):
        put(ws, r, 2 + j, fn(comp), align="center", color=TEAL, bold=True)
    r += 1
band(ws, r, "Package-level", c2=1 + len(comps), fill=SAND, color=TEAL)
r += 1
PKG_ROWS = [("Labor bridge, month over month: lands on Aug compensation", "=" + CHECKS[("labor", "MOM")]),
            ("Labor bridge, month over month: headcount & mix adds across departments", "=" + CHECKS[("labor_hm", "MOM")]),
            ("Labor bridge vs budget: lands on Aug compensation", "=" + CHECKS[("labor", "BUD_M")]),
            ("Labor bridge vs budget: headcount & mix adds across departments", "=" + CHECKS[("labor_hm", "BUD_M")]),
            ("Watch list: net income effect by month equals SMPL's (all views agree)", "=" + CHECKS[("watch",)]),
            ("Watch list: outcome formulas match SMPL's scoring", "=" + CHECKS[("score",)])]
p_first = r
for lbl, f in PKG_ROWS:
    put(ws, r, 1, lbl)
    put(ws, r, 2, f, align="center", color=TEAL, bold=True)
    r += 1
r += 1
put(ws, r, 1, "All tie-outs", bold=True, size=11)
put(ws, r, 2, f'=IF(COUNTIF(B{t_first}:{L(1 + len(comps))}{r - 2},"✗")=0,"✓ All pass","✗ Check the rows above")', bold=True, size=11, color=TEAL)
ALL_TIES = f"'Tie-outs'!B{r}"

# ------------------------------------------------------------------------------------------------
# Cover
# ------------------------------------------------------------------------------------------------
ws = ws_cover
ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 34
for col in "CDEFGH":
    ws.column_dimensions[col].width = 16
put(ws, 1, 1, "SMPL · MD&A variance commentary", bold=True, size=18, color=TEAL)
put(ws, 2, 1, f"{stamp['period']} close · {stamp['entity']} · {stamp['book']} · {stamp['currency']} thousands", size=11, color=GREY)
r = 4
band(ws, r, "Package", c2=8)
r += 1
info = [("Package ID", stamp["packageId"]), ("Status", f"{stamp['status']} · {stamp['approver']['by']} · {stamp['approver']['at']}"),
        ("Data snapshot", stamp["snapshot"]), ("Content hash", stamp["contentHash"]), ("Generated", f"{stamp['generatedAt']} · {stamp['generatedBy']}"),
        ("Tie-outs", None),
        ("If a comparison is reopened", "SMPL marks this package Replaced and issues a new package ID. Check the ID in SMPL before distributing."),
        ("Data", "Synthetic fixture for the SMPL specification; not company data.")]
for k, v in info:
    put(ws, r, 1, k, bold=True)
    if v is None:
        put(ws, r, 2, f"={ALL_TIES}", bold=True, color=TEAL)
    else:
        put(ws, r, 2, v, wrap=True)
        if len(v) > 40:
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
    r += 1
r += 1
band(ws, r, "Comparisons and sign-off", c2=8)
r += 1
for i, h in enumerate(["Tab", "Comparison", "Status", "Approved by", "Approved", "Lines in scope", "Reviewed", "Tie-outs"], 1):
    put(ws, r, i, h, bold=True, fill=TEAL_LIGHT, wrap=True)
r += 1
for j, comp in enumerate(comps):
    ap = comp["approval"]
    c = put(ws, r, 1, comp["tab"], color="0E5A5F")
    c.hyperlink = f"#'{comp['tab']}'!A1"
    c.font = Font(name=FONT, color=TEAL, underline="single")
    put(ws, r, 2, comp["label"], wrap=True)
    put(ws, r, 3, ap["status"], bold=True, color=TEAL)
    put(ws, r, 4, ap["by"])
    put(ws, r, 5, ap["at"])
    put(ws, r, 6, ap["inScope"], align="right")
    put(ws, r, 7, ap["reviewed"], align="right")
    col = L(2 + j)
    put(ws, r, 8, f"=IF(COUNTIF('Tie-outs'!{col}6:{col}40,\"✗\")=0,\"✓\",\"✗\")", align="center", bold=True, color=TEAL)
    r += 1
r += 1
band(ws, r, "Contents", c2=8)
r += 1
CONTENTS = [("Executive summary", "AI-written summary for each comparison, key lines, and the check that cited components add up to the net income change"),
            ("NI walk", "Comparison net income to August net income by statement subtotal, for the accounting and management views"),
            ("Comparison matrix", "Variance of every statement line across all seven comparisons")] + \
           [(c["tab"], f"{c['label']}: income statement, management P&L, GL detail, balance sheet and cash flow, with approved commentary") for c in comps] + \
           [("Labor bridge", "Compensation bridge month over month and vs budget, company and by department (§18)"),
            ("Watch list", "Future impacts by month with the net income effect, and how last month's watch list played out"),
            ("Tie-outs", "Every validation for every comparison, from SMPL and recomputed here"),
            ("Line data", "Every line of every view and comparison in one table, with driver sums"),
            ("Commentary drivers", "The drivers and component lines each paragraph cites")]
for name, desc in CONTENTS:
    c = put(ws, r, 1, name)
    c.hyperlink = f"#'{name}'!A1"
    c.font = Font(name=FONT, color=TEAL, underline="single")
    put(ws, r, 2, desc)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
    r += 1
r += 1
band(ws, r, "How to read this workbook", c2=8)
r += 1
LEGEND = [("Blue numbers", "Ledger amounts from the data snapshot (and the labor engine's bars)."),
          ("Black numbers", "Formulas: roll-ups, variances, percentages, walks, checks."),
          ("Green numbers", "Links to another sheet in this workbook."),
          ("✓ / ✗", "A tie-out that passes or fails. Every ✓ is a live formula except SMPL's export validations."),
          ("F / U", "Favorable or unfavorable for the line (revenue up or cost down is favorable)."),
          ("Commentary", "One paragraph per line. Its amounts are the top three drivers (or component lines) and Other, rounded so they add to the displayed variance."),
          ("Outline", "On each comparison tab, GL lines sit under the statement line built from them; collapse with the outline buttons for a statement-level read.")]
for k, v in LEGEND:
    put(ws, r, 1, k, bold=True, color=BLUE if k == "Blue numbers" else GREEN if k == "Green numbers" else BLACK)
    put(ws, r, 2, v, wrap=True)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
    ws.row_dimensions[r].height = 15 if len(v) < 110 else 28
    r += 1

for w_ in wb.worksheets:
    w_.sheet_view.showGridLines = False if w_.title in ("Cover", "Executive summary") else True
ws_cover.sheet_properties.tabColor = TEAL
ws_exec.sheet_properties.tabColor = TEAL
wb.active = 0
wb.save(OUT)
print("wrote", OUT, "sheets:", [w.title for w in wb.worksheets])
