"""Fill Sage Intacct Marketplace Partner Security Questionnaire from docs/soc2.

Honest / under-claim: pursuing SOC 2 Type I readiness — NOT SOC 2 certified.
"""
from __future__ import annotations

import shutil
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SRC = Path(r"c:\Users\mattj\Downloads\Sage Intacct - Marketplace Partner Security Questionnaire II 2021.xlsx")
OUT_DL = Path(
    r"c:\Users\mattj\Downloads\Sage_Intacct_Marketplace_Partner_Security_Questionnaire_II_2021_SMPL_DRAFT.xlsx"
)
OUT_REPO = Path(
    r"C:\Users\mattj\.cursor\projects\empty-window\saas-financial-intelligence"
    r"\docs\soc2\evidence\questionnaires"
    r"\Sage_Intacct_Marketplace_Partner_Security_Questionnaire_II_2021_SMPL_DRAFT.xlsx"
)
NOTES_MD = OUT_REPO.with_name(OUT_REPO.name.replace(".xlsx", "_NOTES.md"))

DRAFT_DATE = date(2026, 8, 4)


def set_ans(answers: dict, row: int, b, c=None, status="answered", note=""):
    answers[row] = {"b": b, "c": c, "status": status, "note": note}


def build_answers() -> dict:
    answers: dict = {}

    set_ans(answers, 3, "SMPL (SMPL.ai / smpl-ai.com)", None, "answered", "Legal entity name TBD if different from brand")
    set_ans(answers, 4, "Matt Justice", None, "answered")
    set_ans(answers, 5, "Founder; Executive Sponsor & Security Owner", None, "answered")
    set_ans(answers, 6, "mattjustice@smpl-ai.com", None, "answered")
    set_ans(answers, 7, "To be confirmed", None, "needs_matt", "Telephone not in soc2 docs")
    set_ans(answers, 8, DRAFT_DATE.isoformat(), None, "answered", "Draft date; update on final send")

    set_ans(
        answers,
        10,
        (
            "SMPL.ai is a B2B SaaS financial intelligence / FP&A platform. It helps finance teams "
            "read and reconcile finance data (ARR, close, board export). SMPL does not write back "
            "to the customer GL or ERP. Marketplace use with Sage Intacct is for read/reconcile "
            "integration flows."
        ),
        None,
        "answered",
    )

    set_ans(
        answers,
        12,
        "No",
        (
            "Does not store full payment card numbers (SMPL billing via Stripe) or SSNs. May store "
            "business financial data for FP&A (e.g., GL-derived facts, GL account codes, metrics). "
            "No bank-account / PAN storage as payment instruments."
        ),
        "answered",
        "If Sage treats GL account codes as Account Numbers, Matt may prefer Yes + same explanation",
    )

    set_ans(
        answers,
        13,
        "Yes",
        (
            "Auth identifiers (email, session, org membership); multi-tenant financial warehouse "
            "facts (GL-derived, ARR, pipeline, headcount, etc.); export artifacts (Excel close packs, "
            "board decks); billing contact / subscription metadata via Stripe (no full PAN at SMPL); "
            "limited LLM prompt context derived from engine outputs (Anthropic API; keys on backend only)."
        ),
        "answered",
    )

    set_ans(
        answers,
        14,
        "Yes",
        (
            "Customer and auth-related data transit over TLS to/from the web app (Vercel), API "
            "(Railway), datastore (Neon), transactional email (Resend), LLM API (Anthropic), and "
            "billing (Stripe). Customer source systems (e.g., Sage Intacct) may provide data via "
            "customer-authorized import/integration paths."
        ),
        "answered",
    )

    set_ans(
        answers,
        19,
        "Yes",
        (
            "Core information security policies approved internally (P01-P12; P15 AI/LLM). "
            "Approved policies are not SOC 2 certification. SMPL is pursuing SOC 2 Type I readiness "
            "(Security + Availability + Confidentiality) and is NOT SOC 2 certified."
        ),
        "answered",
    )

    set_ans(
        answers,
        20,
        "Yes",
        (
            "Matt Justice — Founder; Executive Sponsor & Security Owner "
            "(solo-founder org; same person currently holds security / eng / ops roles)."
        ),
        "answered",
    )

    set_ans(
        answers,
        21,
        "No",
        (
            "No major security breach documented / known to date. "
            "To be confirmed by Matt if any historical incident should be disclosed."
        ),
        "needs_matt",
        "Confirm no historical breach to disclose",
    )

    set_ans(
        answers,
        22,
        "No",
        (
            "No cyber insurance policy documented in current readiness materials. "
            "To be confirmed by Matt (do not invent coverage amounts)."
        ),
        "needs_matt",
        "Confirm whether cyber insurance exists",
    )

    set_ans(
        answers,
        23,
        "Yes",
        (
            "Formal change management / secure SDLC policy (P05): GitHub PRs, protected main, "
            "deploy to Vercel (FE) / Railway (API)."
        ),
        "answered",
    )

    set_ans(
        answers,
        24,
        "Yes",
        (
            "Approved Incident Response Plan (P04). Tabletop exercised 2026-07-28 "
            "(readiness evidence; not SOC 2 certification)."
        ),
        "answered",
    )

    set_ans(
        answers,
        25,
        "No",
        (
            "Customer-facing privacy policy / customer DPA not yet published as a shareable link. "
            "Confidentiality, retention/deletion, and data-handling policies are approved internally. "
            "Privacy Trust Services criteria deferred for first SOC 2 Type I. "
            "DPA/MSA with counsel in progress — not customer-ready."
        ),
        "answered",
    )

    set_ans(
        answers,
        26,
        "No",
        "Not Privacy Shield certified. (Privacy Shield framework discontinued; not applicable.)",
        "answered",
    )

    set_ans(
        answers,
        27,
        "Yes",
        (
            "Incident Response Plan (P04) covers detection, containment, recovery, and customer "
            "notification when material; includes security/privacy-related incidents."
        ),
        "answered",
    )

    set_ans(
        answers,
        28,
        "No",
        (
            "Formal security awareness training program (planned policy P17) not yet established. "
            "Solo-founder posture; training program TBD."
        ),
        "answered",
    )

    set_ans(
        answers,
        29,
        "No",
        "No formal recurring security awareness training requirement in place yet (P17 not started).",
        "answered",
    )

    set_ans(
        answers,
        30,
        "No",
        (
            "SMPL is NOT SOC 2 certified and has no CPA SOC 2 / ISO 27001 / PCI AoC / SSAE18 report "
            "for SMPL to provide. We are pursuing SOC 2 Type I readiness. Subprocessor (vendor) "
            "SOC 2 reports may be shareable under NDA where collected — not a substitute for "
            "SMPL certification."
        ),
        "answered",
    )

    set_ans(
        answers,
        31,
        "No",
        (
            "No third-party external audit report of SMPL systems within the past two years. "
            "Vendor SOC reports for certain subprocessors collected under NDA for readiness review only."
        ),
        "answered",
    )

    set_ans(
        answers,
        32,
        "Yes",
        (
            "Production is cloud-hosted (Vercel, Railway, Neon). Physical / environmental controls "
            "are provided by those subprocessors; SMPL relies on provider controls and reviews "
            "vendor reports under NDA where available. SMPL does not operate its own data center."
        ),
        "answered",
    )

    set_ans(
        answers,
        33,
        "Yes",
        (
            "Physical access to production servers/network equipment is restricted by cloud "
            "providers (Vercel / Railway / Neon). SMPL personnel do not have physical access to "
            "provider data centers."
        ),
        "answered",
    )

    set_ans(
        answers,
        34,
        "N/A",
        (
            "Solo-founder company today (Matt Justice). No formal background-screening program for "
            "additional employees/contractors documented. To be confirmed / established before "
            "hiring privileged roles."
        ),
        "needs_matt",
        "Confirm if any background screening exists",
    )

    set_ans(
        answers,
        35,
        "Yes",
        (
            "Retention & deletion policy (P08): customer offboarding / deletion requests; delete or "
            "anonymize Confidential warehouse + auth data within agreed window; backups expire per "
            "provider window. Relies on cloud-provider media destruction practices for underlying "
            "infrastructure."
        ),
        "answered",
    )

    set_ans(
        answers,
        36,
        "Yes",
        (
            "In transit: TLS to application and API. At rest: customer data in managed Postgres "
            "(Neon) with provider encryption-at-rest defaults. Production secrets in host "
            "env/secret stores (not in source). Exact cipher suites / key management are "
            "provider-managed — do not invent AES/ISO claims beyond this."
        ),
        "answered",
    )

    set_ans(
        answers,
        37,
        "Yes",
        (
            "Primary datastore backups / PITR via Neon. Backup & restore policy (P12). Restore "
            "fire-drill to throwaway branch completed 2026-07-27 (Pass; production connection "
            "strings unchanged). Laptop/PC backups: recommended full-disk encryption / local "
            "practice — not a formal enterprise backup suite."
        ),
        "answered",
    )

    set_ans(
        answers,
        38,
        "No",
        (
            "Formal Logging & Monitoring policy (P14) not yet published. Operational reliance on "
            "provider consoles/alerts and dependency/secret scanning (Dependabot, GitHub Secret "
            "Protection). No dedicated SIEM/SOC claimed."
        ),
        "answered",
    )

    set_ans(
        answers,
        39,
        "No",
        (
            "No annual third-party penetration test or full internal/external vulnerability "
            "assessment program documented to date. Dependency vulnerability alerting via "
            "Dependabot is enabled (not a substitute for annual app pen tests)."
        ),
        "answered",
    )

    set_ans(
        answers,
        40,
        "Yes",
        (
            "Dependency and secret scanning enabled on the application repository (Dependabot, "
            "GitHub Secret Protection / push protection — confirmed 2026-07-28). Documented risk "
            "assessment (P10). Formal Vulnerability Management policy (P13) not started."
        ),
        "answered",
    )

    set_ans(
        answers,
        41,
        "Yes",
        (
            "Approved Business Continuity / Disaster Recovery policy (P11) and Backup & Restore "
            "policy (P12), honest for managed-cloud solo-founder SaaS. Working RTO/RPO are internal "
            "targets, not published customer SLAs unless MSA states otherwise."
        ),
        "answered",
    )

    set_ans(
        answers,
        42,
        "Yes",
        (
            "Stripe for SMPL subscription billing / checkout. Card data processed by Stripe; SMPL "
            "does not store full PANs. Not used as a customer payment gateway inside Sage Intacct."
        ),
        "answered",
    )

    set_ans(
        answers,
        43,
        "Yes",
        "Secure SDLC / change management (P05): branch -> PR -> protected main -> deploy (Vercel FE / Railway API).",
        "answered",
    )

    set_ans(answers, 44, "Yes", "GitHub for source control and CI.", "answered")

    set_ans(
        answers,
        45,
        "Yes",
        (
            "PR-based code review; Dependabot / secret scanning; tenant-isolation readiness testing "
            "completed 2026-07-29 (Org A != Org B). Scope/frequency: continuous for dependency "
            "alerts; isolation test as readiness evidence. No scheduled third-party application "
            "pen test yet."
        ),
        "answered",
    )

    set_ans(
        answers,
        46,
        "No",
        (
            "No formal OWASP / application-security-specific training program required today "
            "(solo founder; P17 awareness training not started)."
        ),
        "answered",
    )

    set_ans(
        answers,
        47,
        "No",
        "No native mobile application. Responsive web app only (www.smpl-ai.com).",
        "answered",
    )

    set_ans(
        answers,
        48,
        "Yes",
        (
            "Product cloud integrations / subprocessors typically include: Vercel (web), Railway "
            "(API), Neon (Postgres, AWS us-east-1), Resend (email), Anthropic (LLM commentary API), "
            "Stripe (billing), GitHub (source/CI). Customer source systems (e.g., Sage Intacct, "
            "other ERPs/CRMs) are customer-controlled; SMPL reads/reconciles via authorized "
            "import/integration. Marketing CMS (Sanity) and sales CRM (HubSpot) are org tools, "
            "not product Customer Data exhibit."
        ),
        "answered",
    )

    set_ans(
        answers,
        49,
        "Yes",
        (
            "Yes — public cloud / shared services: Vercel, Railway, Neon (managed Postgres on AWS), "
            "GitHub. Development and production use these platforms. Other vendor regions TBD "
            "except Neon us-east-1."
        ),
        "answered",
    )

    set_ans(
        answers,
        51,
        "No",
        (
            "No published customer-facing privacy policy URL to provide yet. Internal "
            "confidentiality / retention / data-handling policies approved. Customer DPA/MSA "
            "legal workstream in progress with counsel — not customer-ready. Privacy Trust "
            "Services deferred for first SOC 2 Type I."
        ),
        "answered",
    )

    set_ans(answers, 52, "No", "Not Privacy Shield certified / listed.", "answered")

    set_ans(
        answers,
        53,
        "Yes",
        (
            "Matt Justice (Founder; Security Owner) currently owns privacy/compliance questions "
            "for the organization (solo-founder)."
        ),
        "assumed",
        "Confirm Matt wants to claim dedicated privacy owner vs N/A given Privacy deferred",
    )

    set_ans(
        answers,
        54,
        "No",
        (
            "No separate regular privacy risk assessment program. Broader information security "
            "risk assessment documented (P10). Privacy TSC deferred for Type I."
        ),
        "answered",
    )

    set_ans(
        answers,
        55,
        "Yes",
        (
            "Incident Response Plan (P04) includes process for security and privacy-related "
            "incidents and customer notification when material."
        ),
        "answered",
    )

    set_ans(
        answers,
        56,
        "No",
        (
            "GDPR DPA / readiness exhibit not yet customer-offerable. Outline sent toward counsel; "
            "firm selection / redline still open. Do not claim GDPR certification."
        ),
        "needs_matt",
        "Confirm DPA status before send; still No unless customer-ready DPA exists",
    )

    set_ans(
        answers,
        57,
        "No",
        (
            "CCPA DPA / readiness exhibit not yet customer-offerable (same legal workstream as "
            "GDPR/DPA). Do not claim CCPA certification."
        ),
        "needs_matt",
        "Same as GDPR DPA status",
    )

    set_ans(
        answers,
        59,
        "No",
        (
            "Not impacted by SolarWinds Sunburst to our knowledge. SMPL does not rely on "
            "SolarWinds Orion for production."
        ),
        "assumed",
        "Confirm Matt has no SolarWinds exposure",
    )

    set_ans(
        answers,
        61,
        "No",
        (
            "Not impacted by Microsoft on-premises Exchange zero-day issues to our knowledge. "
            "SMPL does not operate on-premises Microsoft Exchange."
        ),
        "assumed",
        "Confirm no on-prem Exchange exposure",
    )

    return answers


Q_LABELS = {
    3: "Company name",
    4: "Completer name",
    5: "Job title",
    6: "Email",
    7: "Telephone",
    8: "Date of response",
    10: "App description",
    12: "CC/SSN/Account numbers",
    13: "Storing data?",
    14: "Data in transit?",
    19: "Written infoSec policies",
    20: "Dedicated security person",
    21: "Major breach?",
    22: "Cyber insurance",
    23: "Change control",
    24: "Incident management",
    25: "Privacy policy",
    26: "Privacy Shield",
    27: "Incident response process",
    28: "Security awareness program",
    29: "Awareness training required",
    30: "Compliance certs (SOC/ISO/PCI)",
    31: "Third-party audits (2y)",
    32: "Physical/env controls",
    33: "Physical access restricted",
    34: "Background screening",
    35: "Secure destruction",
    36: "Encryption",
    37: "Backups",
    38: "Security monitoring",
    39: "Annual vuln tests",
    40: "Vuln management process",
    41: "BC/DR policy",
    42: "Payment gateway",
    43: "SDLC",
    44: "Source control",
    45: "App security reviews",
    46: "AppSec training (OWASP)",
    47: "Mobile app",
    48: "Other cloud integrations",
    49: "Public cloud / shared services",
    51: "Privacy policy + link",
    52: "Privacy Shield listed",
    53: "Privacy owner",
    54: "Privacy risk assessments",
    55: "Privacy incident process",
    56: "GDPR DPA",
    57: "CCPA DPA",
    59: "SolarWinds impact",
    61: "MS on-prem zero-day impact",
}

TOP_ITEMS = [
    "1) Confirm NOT claiming SOC 2 certified anywhere in the package (draft already says No / readiness only).",
    "2) Telephone number (row 7) — currently 'To be confirmed'.",
    "3) Cyber insurance (row 22) — answered No; change only if a policy exists (do not invent limits).",
    "4) Major breach history (row 21) — answered No; disclose if anything material exists.",
    "5) GDPR/CCPA DPA (rows 56-57) — answered No; flip only if customer-ready DPA exists.",
    "6) Background screening (row 34) — N/A for solo founder; confirm.",
    "7) Legal entity / company name for Sage contracts if different from 'SMPL / SMPL.ai'.",
    "8) Optional: whether row 12 should be Yes because GL account codes are stored (draft = No + explanation).",
    "9) Privacy owner claim (row 53) — Yes as Matt; switch to N/A if preferred given Privacy deferred.",
    "10) SolarWinds / MS on-prem (rows 59/61) — answered No based on stack; confirm.",
]


def write_notes_sheet(wb, answers: dict, counts: dict):
    if "NOTES" in wb.sheetnames:
        del wb["NOTES"]
    ns = wb.create_sheet("NOTES", 0)
    ns["A1"] = "SMPL DRAFT — Sage Intacct Marketplace Partner Security Questionnaire II 2021"
    ns["A1"].font = Font(bold=True, size=14)
    ns["A2"] = (
        "Status: DRAFT for Matt review. Prefer under-claim. NOT SOC 2 certified. "
        "Do not send until Matt confirms flagged items."
    )
    ns["A2"].font = Font(italic=True, color="C00000")
    ns["A3"] = (
        f"Draft prepared: {DRAFT_DATE.isoformat()} | Sources: docs/soc2/ "
        "(PROGRESS, SECURITY_ONE_PAGER, policies, subprocessors, evidence)"
    )
    ns["A4"] = (
        "Company: SMPL / SMPL.ai (smpl-ai.com) | Completer: Matt Justice | "
        "Product: SaaS FP&A / financial intelligence (read/reconcile; no GL write-back)"
    )

    headers = [
        "Row",
        "Question (short)",
        "Response",
        "Additional (summary)",
        "Status",
        "Note / confirmation needed",
    ]
    for i, h in enumerate(headers, 1):
        cell = ns.cell(6, i, h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E2F3")

    status_fill = {
        "answered": PatternFill("solid", fgColor="C6EFCE"),
        "assumed": PatternFill("solid", fgColor="FFEB9C"),
        "needs_matt": PatternFill("solid", fgColor="FFC7CE"),
    }

    r = 7
    for row in sorted(answers.keys()):
        meta = answers[row]
        ns.cell(r, 1, row)
        ns.cell(r, 2, Q_LABELS.get(row, f"Row {row}"))
        ns.cell(r, 3, meta["b"])
        csum = meta["c"] or ""
        if len(csum) > 180:
            csum = csum[:177] + "..."
        ns.cell(r, 4, csum)
        sc = ns.cell(r, 5, meta["status"])
        sc.fill = status_fill.get(meta["status"], PatternFill())
        ns.cell(r, 6, meta["note"])
        r += 1

    r += 1
    ns.cell(r, 1, "COUNTS").font = Font(bold=True)
    r += 1
    ns.cell(r, 1, "answered (from docs)")
    ns.cell(r, 2, counts["answered"])
    r += 1
    ns.cell(r, 1, "assumed (reasonable; still verify)")
    ns.cell(r, 2, counts["assumed"])
    r += 1
    ns.cell(r, 1, "needs_matt (confirm before send)")
    ns.cell(r, 2, counts["needs_matt"])
    r += 1
    ns.cell(r, 1, "TOTAL questionnaire items filled")
    ns.cell(r, 2, sum(counts.values()))

    r += 2
    ns.cell(r, 1, "TOP ITEMS MATT MUST VERIFY BEFORE SENDING TO SAGE").font = Font(
        bold=True, color="C00000"
    )
    for t in TOP_ITEMS:
        r += 1
        ns.cell(r, 1, t)

    for col, width in enumerate([8, 36, 18, 70, 14, 55], 1):
        ns.column_dimensions[get_column_letter(col)].width = width


def write_notes_md(answers: dict, counts: dict):
    lines = [
        "# Sage Intacct Marketplace Partner Security Questionnaire — SMPL DRAFT NOTES",
        "",
        "**Status:** DRAFT for Matt review. Prefer under-claim. **Not SOC 2 certified.**",
        "",
        f"- **Filled workbook (Downloads):** `{OUT_DL}`",
        f"- **Filled workbook (repo):** `{OUT_REPO}`",
        f"- **Source questionnaire:** `{SRC.name}`",
        f"- **Draft date:** {DRAFT_DATE.isoformat()}",
        "- **Company:** SMPL / SMPL.ai (smpl-ai.com)",
        "- **Completer:** Matt Justice (Founder; Executive Sponsor & Security Owner)",
        "- **Product framing:** SaaS FP&A / financial intelligence — reads/reconciles; **no GL write-back**",
        "- **Compliance framing:** Pursuing SOC 2 Type I readiness (Security + Availability + Confidentiality). **Not certified.**",
        "",
        "## Counts",
        "",
        "| Status | Count |",
        "|--------|------:|",
        f"| Answered from docs | {counts['answered']} |",
        f"| Assumed (verify) | {counts['assumed']} |",
        f"| Needs Matt confirmation | {counts['needs_matt']} |",
        f"| **Total items filled** | **{sum(counts.values())}** |",
        "",
        "## Needs Matt confirmation",
        "",
    ]
    for row in sorted(answers):
        if answers[row]["status"] == "needs_matt":
            lines.append(
                f"- **Row {row} ({Q_LABELS.get(row)}):** Response=`{answers[row]['b']}` — {answers[row]['note']}"
            )
    lines += ["", "## Assumed (still verify)", ""]
    for row in sorted(answers):
        if answers[row]["status"] == "assumed":
            lines.append(
                f"- **Row {row} ({Q_LABELS.get(row)}):** Response=`{answers[row]['b']}` — {answers[row]['note']}"
            )
    lines += ["", "## Top items before sending to Sage", ""]
    lines.extend(TOP_ITEMS)
    lines += [
        "",
        "## Intentionally under-claimed (No / N/A)",
        "",
        "- No SMPL SOC 2 / ISO / PCI / SSAE18 report",
        "- No Privacy Shield",
        "- No formal security awareness / OWASP training program yet (P17 not started)",
        "- No formal logging/monitoring policy yet (P14 not started)",
        "- No annual pen test / vulnerability test program",
        "- No customer-ready GDPR/CCPA DPA yet",
        "- No published privacy policy URL",
        "- No cyber insurance claimed",
        "",
        "## Primary sources used",
        "",
        "- `docs/soc2/PROGRESS.md`",
        "- `docs/soc2/SECURITY_ONE_PAGER.md`",
        "- `docs/soc2/01_system_boundary.md`",
        "- `docs/soc2/02_subprocessors.md`",
        "- `docs/soc2/04_policy_index.md` + policies P01-P12, P15",
        "- Evidence: IR tabletop, neon restore, tenant isolation, access review, Dependabot",
        "",
    ]
    NOTES_MD.parent.mkdir(parents=True, exist_ok=True)
    NOTES_MD.write_text("\n".join(lines), encoding="utf-8")


def main():
    answers = build_answers()
    counts = {"answered": 0, "assumed": 0, "needs_matt": 0}
    for meta in answers.values():
        counts[meta["status"]] = counts.get(meta["status"], 0) + 1

    wb = openpyxl.load_workbook(SRC)
    ws = wb["Marketplace 2021"]
    wrap = Alignment(wrap_text=True, vertical="top")
    for row, meta in answers.items():
        ws.cell(row, 2).value = meta["b"]
        ws.cell(row, 2).alignment = wrap
        if meta["c"] is not None:
            ws.cell(row, 3).value = meta["c"]
            ws.cell(row, 3).alignment = wrap

    write_notes_sheet(wb, answers, counts)
    OUT_DL.parent.mkdir(parents=True, exist_ok=True)
    OUT_REPO.parent.mkdir(parents=True, exist_ok=True)
    # Save once then copy — openpyxl image streams close after first save.
    wb.save(OUT_DL)
    shutil.copy2(OUT_DL, OUT_REPO)
    write_notes_md(answers, counts)

    print("Saved:", OUT_DL)
    print("Saved:", OUT_REPO)
    print("Notes:", NOTES_MD)
    print("COUNTS", counts, "TOTAL", sum(counts.values()))


if __name__ == "__main__":
    main()
