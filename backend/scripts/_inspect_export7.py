"""Compare template vs export (7) text box geometry."""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation


def iter_shapes(shapes):
    for s in shapes:
        yield s
        if hasattr(s, "shapes"):
            yield from iter_shapes(s.shapes)


def shape_text(shape) -> str:
    if not hasattr(shape, "text_frame") or shape.text_frame is None:
        return ""
    return "\n".join(p.text for p in shape.text_frame.paragraphs if p.text).strip()


def para_count(shape) -> int:
    if not hasattr(shape, "text_frame") or shape.text_frame is None:
        return 0
    return len([p for p in shape.text_frame.paragraphs if (p.text or "").strip()])


def inspect_slide(slide, slide_num: int, label: str) -> None:
    print(f"\n=== {label} SLIDE {slide_num} ===")
    kt = None
    for shape in iter_shapes(slide.shapes):
        if shape_text(shape).strip().lower() == "key takeaways":
            kt = shape
            print(f"  KT: top={shape.top} h={shape.height} bottom={shape.top + shape.height}")
    items = []
    for shape in iter_shapes(slide.shapes):
        t = shape_text(shape)
        if not t or len(t) < 25:
            continue
        if "confidential" in t.lower() and "board operating" in t.lower():
            continue
        if t.strip().lower() == "key takeaways":
            continue
        items.append(
            {
                "top": shape.top,
                "h": shape.height,
                "bottom": shape.top + shape.height,
                "len": len(t),
                "paras": para_count(shape),
                "t": t[:100].replace("\n", " | "),
            }
        )
    items.sort(key=lambda x: x["top"])
    for i, it in enumerate(items[:10]):
        ov = ""
        if i + 1 < len(items) and it["bottom"] > items[i + 1]["top"]:
            ov = f" OVERLAP={it['bottom'] - items[i + 1]['top']}"
        gap = it["top"] - (items[i - 1]["bottom"] if i else (kt.top + kt.height if kt else 0))
        print(f"  top={it['top']} h={it['h']} len={it['len']} paras={it['paras']} gap={gap}{ov}")
        print(f"    {it['t']}")


def main() -> None:
    out = Path(r"C:\Users\mattj\Downloads\board_package_2026-06 (7).pptx")
    tmpl = Path(__file__).resolve().parents[1] / "templates/board/SMPL_Board_Review_Template.pptx"
    for path, label in [(tmpl, "TEMPLATE"), (out, "OUTPUT")]:
        prs = Presentation(str(path))
        for sn in [3, 4, 5, 6, 7, 11]:
            if sn == 7:
                print(f"\n=== {label} SLIDE 7 GTM commentary ===")
                for shape in iter_shapes(prs.slides[6].shapes):
                    t = shape_text(shape)
                    if shape.left > 6000000 and len(t) > 15 and "commentary" not in t.lower():
                        print(f"  top={shape.top} len={len(t)}: {t[:75]}")
                continue
            if sn == 11:
                print(f"\n=== {label} SLIDE 11 ===")
                for shape in iter_shapes(prs.slides[10].shapes):
                    t = shape_text(shape)
                    if len(t) > 40 and "confidential" not in t.lower():
                        print(
                            f"  top={shape.top} h={shape.height} len={len(t)} "
                            f"paras={para_count(shape)}: {t[:80].replace(chr(10),' | ')}"
                        )
                continue
            inspect_slide(prs.slides[sn - 1], sn, label)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
