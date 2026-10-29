"""Statement commentary: paragraphs that tie, blank detail, and the compensation sentence."""

from __future__ import annotations

import io
import uuid
from datetime import datetime, timezone
from types import SimpleNamespace
from unittest.mock import MagicMock

import openpyxl

from app.services import validation_engine as ve
from app.services.reporting.export.mda_variance_package import build_mda_package, export_filename
from app.services.reporting.statement_commentary import (
    build_statement_commentary,
    commentary_workbook,
    missing_review_keys,
    present_commentary,
    text_problem,
)


def _books() -> tuple[dict, dict, list]:
    def row(**values):
        return values

    actual_is = {
        "2026-04": row(revenue=90, cogs=30, gross_profit=60, total_opex=20, da=5, interest=2, tax=3, net_income=30, sm=10, rd=6, ga=4),
        "2026-05": row(revenue=80, cogs=20, gross_profit=60, total_opex=30, da=5, interest=2, tax=3, net_income=20, sm=12, rd=10, ga=8),
        "2026-06": row(revenue=100, cogs=30, gross_profit=70, total_opex=40, da=5, interest=2, tax=3, net_income=20, sm=16, rd=14, ga=10),
    }
    # Q1 is Jan-Mar. One month missing on purpose is not this fixture.
    for period, revenue in (("2026-01", 70), ("2026-02", 75), ("2026-03", 85)):
        actual_is[period] = row(
            revenue=revenue, cogs=20, gross_profit=revenue - 20, total_opex=25, da=5, interest=2, tax=3,
            net_income=revenue - 20 - 25 - 5 - 2 - 3, sm=10, rd=8, ga=7,
        )
    actual_bs = {
        "2026-06": row(cash=30, ar=10, total_assets=80, total_liabilities=40, total_le=80),
        "2026-05": row(cash=28, ar=9, total_assets=76, total_liabilities=38, total_le=76),
    }
    actual_cfs = {
        "2026-06": row(cfo=12, cfi=-4, cff=2, net_change=10, ending_cash=30),
        "2026-05": row(cfo=8, cfi=-2, cff=1, net_change=7, ending_cash=28),
    }
    ts = {
        "Actual": {
            "is": actual_is,
            "bs": actual_bs,
            "cfs": actual_cfs,
            "gl_opex": {
                "2026-06": {"sm|Sales|Salaries": 20_000, "sm|Sales|Travel": 5},
                "2026-05": {"sm|Sales|Salaries": 8_000},
            },
        },
        "Budget": {
            "is": {"2026-06": row(revenue=90, cogs=25, gross_profit=65, total_opex=35, da=5, interest=2, tax=3, net_income=20)},
            "bs": {"2026-06": row(cash=27, total_assets=70, total_liabilities=35, total_le=70)},
            "cfs": {"2026-06": row(cfo=9, cfi=-3, cff=1, net_change=7, ending_cash=27)},
            "gl_opex": {"2026-06": {"sm|Sales|Salaries": 15_000, "ga|Finance|Rent": 1_000}},
        },
        "Forecast": {"is": {}, "bs": {}, "cfs": {}},
    }
    history = {
        "is": {"2025-06": row(revenue=50, cogs=15, gross_profit=35, total_opex=20, da=4, interest=1, tax=2, net_income=8)},
        "bs": {"2025-06": row(cash=10, total_assets=40, total_liabilities=20, total_le=40)},
        "cfs": {"2025-06": row(cfo=4, cfi=-1, cff=0, net_change=3, ending_cash=10)},
    }
    bridge = {
        "id": "mom",
        "available": True,
        "from": {"total": 100},
        "to": {"total": 180},
        "bars": [
            {"id": "hires", "amount": 50},
            {"id": "departures", "amount": -20},
            {"id": "rate", "amount": 10},
            {"id": "reconciling", "amount": 40},
        ],
    }
    return ts, history, [bridge]


def _comparison(package, comparison_id):
    return next(item for item in package["comparisons"] if item["id"] == comparison_id)


def _statement(package, comparison_id, statement_id):
    return next(item for item in _comparison(package, comparison_id)["statements"] if item["id"] == statement_id)


def _line(statement, line_id):
    return next(item for item in statement["lines"] if item["id"] == line_id)


def test_mda_package_uses_balances_and_whole_dollars() -> None:
    ts, history, bridges = _books()
    package = build_statement_commentary(close="2026-06", ts=ts, history=history, labor_bridges=bridges)
    built = build_mda_package(package, ts=ts, history=history)
    assert built["stamp"]["status"] == "Draft"
    assert built["stamp"]["precision"]["divisor"] == 100
    assert built["labor"] == []
    assert export_filename(built).endswith("_DRAFT.xlsx")
    mom = next(item for item in built["comparisons"] if item["key"] == "MOM")
    income = {row["nid"]: row for row in mom["views"]["IS"]["rows"]}
    gross = income["gross_profit"]
    assert gross["fu"] == "F"
    assert "rose $10" in gross["text"]
    assert "higher revenue (+$20)" in gross["text"]
    assert "higher cost of revenue" in gross["text"]
    cost = income["cogs"]
    assert cost["fu"] == "U"
    assert income["interest"]["fu"] == ""
    net = income["net_income"]
    explained = sum(step["step"] for step in mom["walks"]["IS"]["steps"])
    assert explained == net["var"]
    qoq = next(item for item in built["comparisons"] if item["key"] == "QOQ")
    cash = next(row for row in qoq["views"]["BS"]["rows"] if row["nid"] == "cash")
    ending = next(row for row in qoq["views"]["CF"]["rows"] if row["nid"] == "ending_cash")
    assert cash["cur"] == 3000
    assert ending["cur"] == 3000
    assert cash["cmp"] is None


def test_accepted_analysis_is_the_workbook_line():
    ts, history, bridges = _books()
    package = build_statement_commentary(close="2026-06", ts=ts, history=history, labor_bridges=bridges)
    built = build_mda_package(
        package,
        ts=ts,
        history=history,
        analyses=[
            {
                "comparison": "mom",
                "section_id": "cogs",
                "paragraph": "Cost of revenue increased $7.40m. Hosting · Compute rose, and Support · Tools rose $4.25k.",
            },
            {"comparison": "mom", "section_id": "not-a-section", "paragraph": "Ignore this."},
        ],
    )
    mom = next(item for item in built["comparisons"] if item["key"] == "MOM")
    income = {row["nid"]: row for row in mom["views"]["IS"]["rows"]}
    text = income["cogs"]["text"]
    assert text == "Cost of revenue increased $7.40m. Hosting · Compute rose, and Support · Tools rose $4.25k."
    assert "rose $10" in income["gross_profit"]["text"]
    assert "Ignore this." not in text
    yoy = next(item for item in built["comparisons"] if item["key"] == "YOY")
    yoy_cash = next(row for row in yoy["views"]["BS"]["rows"] if row["nid"] == "cash")
    assert yoy_cash["cur"] == 3000
    assert yoy_cash["cmp"] == 1000


def test_ebitda_and_pretax_name_the_lines_above_them():
    ts, history, bridges = _books()
    ts["Actual"]["is"]["2026-05"].update(ebitda=30, pretax=23)
    ts["Actual"]["is"]["2026-06"].update(gross_profit=80, ebitda=40, da=6, interest=3, pretax=31)
    package = build_statement_commentary(close="2026-06", ts=ts, history=history, labor_bridges=bridges)
    income = _statement(package, "mom", "income_statement")
    ebitda = _line(income, "ebitda")
    pretax = _line(income, "pretax")
    assert "gross profit" in ebitda["commentary"]
    assert "operating expenses" in ebitda["commentary"]
    assert "operating income" in pretax["commentary"]
    assert "interest" in pretax["commentary"]
    built = build_mda_package(package, ts=ts, history=history)
    mom = next(item for item in built["comparisons"] if item["key"] == "MOM")
    rows = {row["nid"]: row for row in mom["views"]["IS"]["rows"]}
    assert "higher gross profit" in rows["ebitda"]["text"]
    assert "operating expenses" in rows["ebitda"]["text"]
    assert "operating income" in rows["pretax"]["text"]
    assert "interest" in rows["pretax"]["text"]


def test_month_story_ties_and_compensation_stays_at_three_parts() -> None:
    ts, history, bridges = _books()
    package = build_statement_commentary(close="2026-06", ts=ts, history=history, labor_bridges=bridges)
    income = _statement(package, "mom", "income_statement")
    assert income["ties"] is True
    compensation = _line(income, "compensation")
    text = compensation["commentary"].lower()
    assert "hires $50.00" in text
    assert "departures -$20.00" in text
    assert "everything else $50.00" in text
    assert "rate" not in text
    assert "mix" not in text
    assert income["paragraph"].endswith(compensation["commentary"])
    overview = _comparison(package, "mom")["overview"]
    assert "Income" not in overview
    assert income["paragraph"] in overview
    assert _statement(package, "mom", "balance_sheet")["paragraph"] in overview
    assert _statement(package, "mom", "cash_flow")["paragraph"] in overview


def test_account_detail_stays_blank_when_the_other_side_is_missing() -> None:
    ts, history, bridges = _books()
    package = build_statement_commentary(close="2026-06", ts=ts, history=history, labor_bridges=bridges)
    income = _statement(package, "mom", "income_statement")
    travel = next(row for row in income["accounts"] if row["label"] == "Travel")
    assert travel["comparison_amount"] is None
    assert travel["variance"] is None
    assert travel["review_required"] is False
    sales = next(row for row in income["departments"] if row["label"] == "Sales")
    assert sales["variance"] is None
    assert sales["commentary"] == "Department comparison is incomplete."
    salaries = next(row for row in income["accounts"] if row["label"] == "Salaries")
    assert salaries["variance"] == 12_000
    assert salaries["review_required"] is True
    yoy = _statement(package, "yoy", "income_statement")
    assert yoy["accounts"] == []
    assert yoy["accounts_note"] == "Account and department detail is not used on this comparison."
    assert _line(yoy, "revenue")["variance"] == 50
    forecast = _comparison(package, "forecast")
    assert forecast["available"] is False
    assert forecast["note"] == "No forecast for June 2026."
    assert forecast["statements"] == []


def test_budget_account_needs_both_sides_and_small_lines_do_not_require_review() -> None:
    ts, history, bridges = _books()
    package = build_statement_commentary(close="2026-06", ts=ts, history=history, labor_bridges=bridges)
    income = _statement(package, "budget", "income_statement")
    rent = next(row for row in income["accounts"] if row["label"] == "Rent")
    assert rent["current"] is None
    assert rent["variance"] is None
    assert rent["commentary"] == "No comparison amount."
    revenue = _line(income, "revenue")
    assert revenue["review_required"] is True
    assert revenue["variance"] == 10
    assert _line(income, "sm")["review_required"] is False


def test_quarter_balances_stay_at_period_end_while_flows_sum() -> None:
    ts, history, bridges = _books()
    ts["Actual"]["bs"]["2026-04"] = {"cash": 26, "ar": 8, "total_assets": 70, "total_liabilities": 36, "total_le": 70}
    ts["Actual"]["bs"]["2026-03"] = {"cash": 20, "ar": 7, "total_assets": 60, "total_liabilities": 30, "total_le": 60}
    ts["Actual"]["cfs"]["2026-04"] = {"cfo": 5, "cfi": -1, "cff": 0, "net_change": 4, "ending_cash": 26}
    ts["Actual"]["cfs"]["2026-03"] = {"cfo": 3, "cfi": -1, "cff": 0, "net_change": 2, "ending_cash": 20}
    package = build_statement_commentary(close="2026-06", ts=ts, history=history, labor_bridges=bridges)
    cash = _line(_statement(package, "qoq", "balance_sheet"), "cash")
    assert cash["current"] == 30
    assert cash["comparison_amount"] == 20
    assert cash["variance"] == 10
    receivables = _line(_statement(package, "qoq", "balance_sheet"), "ar")
    assert receivables["current"] == 10
    flows = _statement(package, "qoq", "cash_flow")
    assert _line(flows, "cfo")["current"] == 25
    assert _line(flows, "ending_cash")["current"] == 30
    assert _line(flows, "ending_cash")["comparison_amount"] == 20
    ytd_cash = _line(_statement(package, "ytd", "balance_sheet"), "cash")
    assert ytd_cash["current"] == 30
    assert ytd_cash["comparison_amount"] == 10


def test_quarter_sums_the_months_and_does_not_invent_accounts() -> None:
    ts, history, bridges = _books()
    package = build_statement_commentary(close="2026-06", ts=ts, history=history, labor_bridges=bridges)
    qoq = _statement(package, "qoq", "income_statement")
    # Q2 revenue 90+80+100 = 270. Q1 70+75+85 = 230.
    assert _line(qoq, "revenue")["variance"] == 40
    assert "because" in _line(qoq, "net_income")["commentary"]
    assert "revenue" in _line(qoq, "net_income")["commentary"]
    assert qoq["accounts"] == []
    assert all(line["id"] != "compensation" for line in qoq["lines"])


def test_missing_quarter_month_is_blank_not_zero() -> None:
    ts, history, bridges = _books()
    del ts["Actual"]["is"]["2026-04"]
    package = build_statement_commentary(close="2026-06", ts=ts, history=history, labor_bridges=bridges)
    revenue = _line(_statement(package, "qoq", "income_statement"), "revenue")
    assert revenue["current"] is None
    assert revenue["variance"] is None
    assert "no comparison" in revenue["commentary"].lower()


def test_broken_net_income_walk_does_not_claim_a_tie() -> None:
    ts, history, bridges = _books()
    ts["Actual"]["is"]["2026-06"]["net_income"] = 99
    package = build_statement_commentary(close="2026-06", ts=ts, history=history, labor_bridges=bridges)
    income = _statement(package, "mom", "income_statement")
    assert income["ties"] is False
    assert "do not equal" in income["tie_note"]


def test_unavailable_bridge_states_the_reason_and_stops() -> None:
    ts, history, _bridges = _books()
    bridges = [{"id": "mom", "available": False, "reason": "No GL compensation for May 2026."}]
    package = build_statement_commentary(close="2026-06", ts=ts, history=history, labor_bridges=bridges)
    compensation = _line(_statement(package, "mom", "income_statement"), "compensation")
    assert compensation["commentary"] == "No GL compensation for May 2026."
    assert "$" not in compensation["commentary"]
    assert compensation["review_required"] is True


def test_edits_flow_into_the_paragraph_and_a_new_number_fails() -> None:
    ts, history, bridges = _books()
    package = build_statement_commentary(close="2026-06", ts=ts, history=history, labor_bridges=bridges)
    revenue = _line(_statement(package, "mom", "income_statement"), "revenue")
    assert text_problem(package, {revenue["key"]: "Revenue rose $20.00."}, {}) is None
    assert text_problem(package, {revenue["key"]: "Revenue rose $7.00."}, {}) == "line_does_not_tie"
    assert text_problem(package, {}, {"mom": "Collections were $123.45."}) == "addition_does_not_tie"
    assert text_problem(package, {}, {"mom": "Revenue rose $20.00."}) is None
    presented = present_commentary(package, {"edits": {revenue["key"]: "Revenue rose $20.00."}})
    paragraph = _statement(presented, "mom", "income_statement")["paragraph"]
    assert "Revenue rose $20.00." in paragraph


def test_workbook_leaves_the_blank_forecast_blank() -> None:
    ts, history, bridges = _books()
    package = build_statement_commentary(close="2026-06", ts=ts, history=history, labor_bridges=bridges)
    raw = commentary_workbook(package, allowed=False, ts=ts, history=history)
    book = openpyxl.load_workbook(io.BytesIO(raw))
    assert book["Cover"]["A1"].value == "SMPL Demo Co · MD&A variance commentary"
    assert book["Cover"]["B6"].value == "Draft"
    assert "Executive summary" in book.sheetnames
    assert "NI walk" in book.sheetnames
    assert "MoM" in book.sheetnames
    assert "Labor bridge" not in book.sheetnames
    assert "Forecast M" not in book.sheetnames
    cover = [book["Cover"].cell(row, 1).value for row in range(1, book["Cover"].max_row + 1)]
    assert "Forecast M" in cover
    assert "Labor bridge" in cover
    mom = book["MoM"]
    formulas = [mom.cell(row, 4).value for row in range(1, mom.max_row + 1)]
    assert any(isinstance(value, str) and value.startswith("=B") for value in formulas)


def test_a_wrapper_named_like_the_statement_line_is_not_repeated() -> None:
    from app.services.reporting.export.mda_variance_package import _hoist_same_name_wrapper

    wrapper = {
        "nid": "wrap",
        "cur": 10000,
        "cmp": 6000,
        "row": {"label": "Revenue"},
        "child_nodes": [
            {"nid": "sub", "cur": 6000, "cmp": 4000, "row": {"label": "Subscription"}},
            {"nid": "svc", "cur": 4000, "cmp": 2000, "row": {"label": "Services"}},
        ],
    }
    hoisted = _hoist_same_name_wrapper([wrapper], "Revenue", 10000, 6000)
    assert [node["nid"] for node in hoisted] == ["sub", "svc"]


def test_a_parent_whose_children_sum_can_be_exported() -> None:
    from app.services.reporting.export.mda_variance_package import _detail_nodes

    nodes = _detail_nodes(
        [
            {"id": "ga", "label": "G&A", "depth": 1, "current": 100, "comparison_amount": 80, "line": "ga"},
            {"id": "software", "label": "Software Tools", "depth": 2, "current": 60, "comparison_amount": 50, "line": "ga"},
            {"id": "travel", "label": "Travel", "depth": 2, "current": 40, "comparison_amount": 30, "line": "ga"},
        ],
        "ga",
        {},
        "2026-06-01",
        "May",
    )
    parent = nodes[0]
    assert parent["row"]["kind"] == "sum"
    assert parent["child_nodes"][0]["sentence"]
    assert "Software" in parent["row"]["text"]


def test_workbook_writes_the_rollup_instead_of_the_flat_account_list() -> None:
    from app.services.reporting.gl_income_statement import OPEX_LINE_BY_DEPARTMENT
    from app.services.reporting.gl_rollup import trees_by_period

    ts, history, bridges = _books()
    ts["Actual"]["gl_rollup"] = trees_by_period(
        [
            {"period": "2026-06", "statement_category": "Operating Expense", "account_name": "Software Tools", "department": "Finance", "expense_type": "Software", "amount": 100},
            {"period": "2026-06", "statement_category": "Operating Expense", "account_name": "Software Tools", "department": "G&A", "expense_type": "Software", "amount": 40},
            {"period": "2026-05", "statement_category": "Operating Expense", "account_name": "Software Tools", "department": "Finance", "expense_type": "Software", "amount": 80},
            {"period": "2026-05", "statement_category": "Operating Expense", "account_name": "Software Tools", "department": "G&A", "expense_type": "Software", "amount": 40},
            {"period": "2026-06", "statement_category": "Cost of Revenue", "account_name": "Cloud Hosting", "department": "Engineering", "expense_type": "Cloud Hosting", "amount": 50},
            {"period": "2026-05", "statement_category": "Cost of Revenue", "account_name": "Cloud Hosting", "department": "Engineering", "expense_type": "Cloud Hosting", "amount": 20},
        ],
        department_aliases=OPEX_LINE_BY_DEPARTMENT,
    )
    package = build_statement_commentary(close="2026-06", ts=ts, history=history, labor_bridges=bridges)
    income = _statement(package, "mom", "income_statement")
    labels = [row["label"] for row in income["detail_rows"]]
    assert any(label.startswith("Finance") for label in labels)
    assert any(label.startswith("G&A") for label in labels)
    assert any("Cloud Hosting" in label for label in labels)
    finance = next(row for row in income["detail_rows"] if row["label"].startswith("Finance"))
    general = next(row for row in income["detail_rows"] if row["label"].startswith("G&A"))
    assert finance["variance"] == 20
    assert general["variance"] == 0
    assert finance["current"] == 100
    assert general["current"] == 40
    assert _statement(package, "yoy", "income_statement")["detail_rows"] == []
    book = openpyxl.load_workbook(io.BytesIO(commentary_workbook(package, ts=ts, history=history)))
    shown = [book["MoM"].cell(row, 1).value for row in range(1, book["MoM"].max_row + 1)]
    assert any(isinstance(value, str) and value.strip().startswith("Finance") for value in shown)
    assert any(isinstance(value, str) and "Cloud Hosting" in value for value in shown)
    assert "Salaries" not in shown
    assert "Travel" not in shown


def test_deck_rollup_keeps_depth_one_lines_and_material_children() -> None:
    from app.services.reporting.gl_income_statement import OPEX_LINE_BY_DEPARTMENT
    from app.services.reporting.gl_rollup import trees_by_period
    from app.services.reporting.statement_commentary import income_rollup_lines

    rows = [
        {"period": "2026-06", "statement_category": "Operating Expense", "account_name": "Software Tools", "department": "Finance", "expense_type": "Software", "amount": 100},
        {"period": "2026-06", "statement_category": "Operating Expense", "account_name": "Software Tools", "department": "G&A", "expense_type": "Software", "amount": 40},
        {"period": "2026-05", "statement_category": "Operating Expense", "account_name": "Software Tools", "department": "Finance", "expense_type": "Software", "amount": 80},
        {"period": "2026-05", "statement_category": "Operating Expense", "account_name": "Software Tools", "department": "G&A", "expense_type": "Software", "amount": 40},
        {"period": "2026-06", "statement_category": "Cost of Revenue", "account_name": "Cloud Hosting", "department": "Engineering", "expense_type": "Cloud Hosting", "amount": 50},
        {"period": "2026-05", "statement_category": "Cost of Revenue", "account_name": "Cloud Hosting", "department": "Engineering", "expense_type": "Cloud Hosting", "amount": 20},
        {"period": "2026-06", "statement_category": "Operating Expense", "account_name": "Travel and Entertainment", "department": "Sales", "expense_type": "Travel", "vendor": "Uber", "amount": 20_000},
        {"period": "2026-06", "statement_category": "Operating Expense", "account_name": "Travel and Entertainment", "department": "Sales", "expense_type": "Travel", "vendor": "Cafe", "amount": 5},
        {"period": "2026-05", "statement_category": "Operating Expense", "account_name": "Travel and Entertainment", "department": "Sales", "expense_type": "Travel", "vendor": "Uber", "amount": 1_000},
        {"period": "2026-05", "statement_category": "Operating Expense", "account_name": "Travel and Entertainment", "department": "Sales", "expense_type": "Travel", "vendor": "Cafe", "amount": 5},
    ]
    ts = {"Actual": {"is": {}, "bs": {}, "cfs": {}, "gl_rollup": trees_by_period(rows, department_aliases=OPEX_LINE_BY_DEPARTMENT)}}
    rollup = income_rollup_lines(ts, close="2026-06")
    labels = [row["label"] for row in rollup["mom"]["lines"]]
    assert any(label.startswith("Finance") for label in labels)
    assert any(label.startswith("G&A") for label in labels)
    assert any("Cloud Hosting" in label for label in labels)
    assert any(label == "Uber" for label in labels)
    assert all(label != "Cafe" for label in labels)
    assert rollup["budget"]["lines"] == []
    assert income_rollup_lines(None, close="2026-06")["mom"]["lines"] == []


def _blob(sections: dict | None = None) -> SimpleNamespace:
    return SimpleNamespace(
        status="COMPLETE",
        sections_json=dict(sections or {"kpis": True}),
        built_at=datetime.now(timezone.utc),
        updated_at=None,
    )


def test_allow_waits_for_the_required_line_reviews() -> None:
    org_id = uuid.uuid4()
    blob = _blob()
    db = MagicMock()
    db.scalars.return_value.first.return_value = blob
    package = {
        "comparisons": [
            {
                "id": "mom",
                "available": True,
                "statements": [
                    {
                        "lines": [
                            {"key": "mom|income_statement|revenue", "review_required": True},
                            {"key": "mom|income_statement|sm", "review_required": False},
                        ],
                        "departments": [],
                        "accounts": [{"key": "mom|income_statement|Sales|Travel", "review_required": False}],
                    }
                ],
            }
        ]
    }
    try:
        ve.record_allow(db, org_id, "2026-06", allowed_by="Owner", commentary_package=package)
        raise AssertionError("allow should wait")
    except ValueError as exc:
        assert str(exc) == "commentary_review_open"
    ve.save_commentary(
        db,
        org_id,
        "2026-06",
        reviews={"mom|income_statement|revenue": True},
        edits={},
        additions={},
        reviewed_by="Owner",
        package=package,
    )
    allowed = ve.record_allow(db, org_id, "2026-06", allowed_by="Owner", commentary_package=package)
    assert allowed["allowed"] is True
    assert missing_review_keys(package, blob.sections_json["commentary"]) == []
