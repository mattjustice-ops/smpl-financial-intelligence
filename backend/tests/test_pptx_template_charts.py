"""Phase 3 — live chart/KPI injection into board PPTX template."""

from __future__ import annotations

from decimal import Decimal

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.services.dashboard.schemas import ExecutiveFlowResponse, WaterfallSummaryRow
from app.services.reporting.export.pptx_template_charts import (
    apply_template_charts,
    apply_template_visuals,
    build_template_chart_series,
)
from app.services.reporting.export.pptx_template_export import (
    map_template_slides_to_slide_keys,
    resolve_board_pptx_template,
)
from app.services.reporting.export.schemas import ExportValidationSummary, ReportingBundle


def _wf_row(period: str, wtype: str, amount: Decimal, scenario: str = "Actual") -> WaterfallSummaryRow:
    return WaterfallSummaryRow(
        organization_id="test-org",
        scenario=scenario,
        period=period,
        waterfall_name="arr" if "arr" in wtype or wtype.startswith("new") else "cash_flow",
        waterfall_type=wtype,
        line_item=wtype,
        line_item_order=0,
        amount=amount,
        source_table="test",
    )


def _bundle_with_monthly_series() -> ReportingBundle:
    arr_rows: list[WaterfallSummaryRow] = []
    cash_rows: list[WaterfallSummaryRow] = []
    for m, nn_a, nn_b, cash_a, cash_b in [
        (1, 1_300_000, 1_200_000, 48_000_000, 12_000_000),
        (2, 1_500_000, 1_400_000, 51_000_000, 15_000_000),
        (3, 1_700_000, 1_600_000, 58_000_000, 22_000_000),
        (4, 1_900_000, 1_800_000, 62_000_000, 26_000_000),
        (5, 2_100_000, 1_900_000, 66_000_000, 30_000_000),
        (6, 2_700_000, 2_200_000, 71_000_000, 48_000_000),
    ]:
        p = f"2026-{m:02d}"
        arr_rows.extend(
            [
                _wf_row(p, "new_arr", Decimal(nn_a), "Actual"),
                _wf_row(p, "new_arr", Decimal(nn_b), "Budget"),
                _wf_row(p, "ending_arr", Decimal(80_000_000 + m * 1_000_000), "Actual"),
                _wf_row(p, "ending_arr", Decimal(79_000_000 + m * 1_000_000), "Budget"),
            ]
        )
        cash_rows.extend(
            [
                _wf_row(p, "ending_cash", Decimal(cash_a), "Actual"),
                _wf_row(p, "ending_cash", Decimal(cash_b), "Budget"),
            ]
        )

    return ReportingBundle(
        organization_id="test-org",
        organization_name="Test Co",
        scenario="Combined",
        period_label="June 2026",
        as_of_period="2026-06",
        start_period="2026-01",
        end_period="2026-12",
        currency="USD",
        executive_flow=ExecutiveFlowResponse(
            organization_id="test-org",
            scenario="Combined",
            start_period="2026-01",
            end_period="2026-12",
            as_of_period="2026-06",
        ),
        comparison_waterfalls={"arr": arr_rows, "cash_flow": cash_rows},
        validation=ExportValidationSummary(status="pass"),
    )


def test_build_template_chart_series_arr_and_cash_in_millions():
    bundle = _bundle_with_monthly_series()
    arr = build_template_chart_series(bundle, "arr_net_new")
    assert arr is not None
    cats, series = arr
    assert cats == ["Jan", "Feb", "Mar", "Apr", "May", "Jun"]
    assert series["Net New ARR (Act)"][-1] == 2.7
    assert series["Net New ARR (Bud)"][-1] == 2.2

    cash = build_template_chart_series(bundle, "cash_ending")
    assert cash is not None
    _, cash_series = cash
    assert cash_series["Ending Cash (Actual)"][-1] == 71.0
    assert cash_series["Ending Cash (Budget)"][-1] == 48.0


def test_apply_template_charts_replaces_gold_deck_series():
    template = resolve_board_pptx_template()
    if template is None:
        return

    bundle = _bundle_with_monthly_series()
    prs = Presentation(str(template))
    key_to_idx = map_template_slides_to_slide_keys(prs)
    assert "arr_waterfall" in key_to_idx
    assert "cash_forecast" in key_to_idx

    changed = apply_template_charts(prs, bundle, key_to_idx)
    assert changed >= 2

    arr_slide = prs.slides[key_to_idx["arr_waterfall"] - 1]
    chart = next(s.chart for s in arr_slide.shapes if s.shape_type == MSO_SHAPE_TYPE.CHART)
    cats = [str(c.label) for c in chart.plots[0].categories]
    assert cats == ["Jan", "Feb", "Mar", "Apr", "May", "Jun"]
    act = next(s for s in chart.plots[0].series if "Act" in str(s.name))
    assert round(list(act.values)[-1], 2) == 2.7

    cash_slide = prs.slides[key_to_idx["cash_forecast"] - 1]
    cash_chart = next(s.chart for s in cash_slide.shapes if s.shape_type == MSO_SHAPE_TYPE.CHART)
    cash_act = next(s for s in cash_chart.plots[0].series if "Actual" in str(s.name))
    assert round(list(cash_act.values)[-1], 2) == 71.0


def test_apply_template_visuals_returns_counts():
    template = resolve_board_pptx_template()
    if template is None:
        return

    bundle = _bundle_with_monthly_series()
    prs = Presentation(str(template))
    counts = apply_template_visuals(prs, bundle)
    assert counts["charts"] >= 2
    # Scorecard/KPI cells depend on layout match; charts are the hard dual-ink gate.
    assert isinstance(counts["scorecard"], int)
    assert isinstance(counts["kpis"], int)
