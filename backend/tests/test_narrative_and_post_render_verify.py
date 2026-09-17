"""Non-numeric commentary + post-render deck fidelity checks."""

from __future__ import annotations

import io
from decimal import Decimal

from pptx import Presentation
from pptx.util import Inches, Pt

from app.services.commentary.narrative_verify import (
    soft_strip_filler_and_placeholders,
    verify_narrative_text,
    verify_nested_commentary_narrative,
)
from app.services.reporting.export.deck_post_render_verify import (
    extract_pptx_texts,
    verify_rendered_deck,
)


def _mini_pptx(*texts: str) -> bytes:
    prs = Presentation()
    # Use a blank layout
    layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    for i, text in enumerate(texts):
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4 + i * 0.6), Inches(8), Inches(0.5))
        tf = box.text_frame
        tf.paragraphs[0].text = text
        tf.paragraphs[0].font.size = Pt(12)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_forbidden_filler_detected() -> None:
    result = verify_narrative_text("ARR saw a significant beat versus budget.")
    assert not result.ok
    assert any(i.kind == "filler" for i in result.issues)


def test_placeholder_detected_and_stripped() -> None:
    text = "Primary driver was expansion. TBD more detail next month."
    result = verify_narrative_text(text)
    assert any(i.kind == "placeholder" for i in result.issues)
    cleaned = soft_strip_filler_and_placeholders(text)
    assert "TBD" not in cleaned
    assert "expansion" in cleaned.lower()


def test_direction_beat_with_negative_variance() -> None:
    result = verify_narrative_text(
        "Revenue beat budget by -$370.0K in the close month."
    )
    assert any(i.kind == "direction" for i in result.issues)


def test_direction_miss_with_positive_variance() -> None:
    result = verify_narrative_text(
        "ARR missed plan by +$413.1K versus budget."
    )
    assert any(i.kind == "direction" for i in result.issues)


def test_direction_consistent_ok() -> None:
    result = verify_narrative_text(
        "ARR beat budget by +$413.1K on expansion strength."
    )
    assert not any(i.kind == "direction" for i in result.issues)


def test_scenario_mislabel_closed_month_as_forecast() -> None:
    result = verify_narrative_text(
        "June Forecast cash ends at $70.6M.",
        close_period="2026-06",
    )
    assert any(i.kind == "scenario" for i in result.issues)


def test_scenario_ok_when_actual_labeled() -> None:
    result = verify_narrative_text(
        "June Actual cash ends at $70.6M; H2 Forecast remains open.",
        close_period="2026-06",
    )
    assert not any(i.kind == "scenario" for i in result.issues)


def test_nested_commentary_strips_filler() -> None:
    commentary = {
        "variance_commentary": {
            "Revenue": {
                "what_changed": "It is worth noting that revenue beat budget by +$0.4M.",
                "driver": "Expansion ARR",
            }
        }
    }
    rewritten, result = verify_nested_commentary_narrative(
        commentary, close_period="2026-06", strip=True
    )
    assert not result.ok
    cell = rewritten["variance_commentary"]["Revenue"]["what_changed"]
    assert "worth noting" not in cell.lower()
    assert "Expansion ARR" == rewritten["variance_commentary"]["Revenue"]["driver"]


def test_extract_pptx_texts_roundtrip() -> None:
    raw = _mini_pptx("$86.1M", "Key takeaway: expansion drove net new ARR.")
    slides, texts = extract_pptx_texts(raw)
    assert slides == 1
    assert any("$86.1M" in t for t in texts)
    assert any("expansion" in t.lower() for t in texts)


def test_post_render_anchor_pass() -> None:
    raw = _mini_pptx("$86.1M", "$7.41M", "$70.61M")
    evidence = {
        "period_matrix.ARR.cm.actual": Decimal("86100000"),
        "period_matrix.Revenue.cm.actual": Decimal("7410000"),
        "period_matrix.Cash.cm.actual": Decimal("70610000"),
    }
    result = verify_rendered_deck(raw, evidence=evidence, payload={"close_period": "2026-06"})
    assert result.anchors.ok
    assert result.slide_count == 1


def test_post_render_anchor_fail_soft() -> None:
    raw = _mini_pptx("$1.00M", "Narrative only — no board anchors.")
    evidence = {
        "period_matrix.ARR.cm.actual": Decimal("86100000"),
        "period_matrix.Cash.cm.actual": Decimal("70610000"),
    }
    result = verify_rendered_deck(raw, evidence=evidence, fail_closed=False)
    assert not result.anchors.ok
    assert "anchors:" in result.summary()


def test_post_render_skips_chart_only_prior_month_arr_anchors() -> None:
    """Jan–Apr story.ending_arr keys are chart-series only — do not false-fail anchors."""
    from app.services.reporting.export.deck_post_render_verify import _select_anchors

    evidence = {
        "story.actual.ending_arr.2026-01": Decimal("76310000"),
        "story.actual.ending_arr.2026-02": Decimal("77820000"),
        "story.actual.ending_arr.2026-03": Decimal("79505000"),
        "story.actual.ending_arr.2026-04": Decimal("81385000"),
        "period_matrix.ARR.cm.actual": Decimal("86100000"),
        "period_matrix.Cash.cm.actual": Decimal("70610000"),
    }
    anchors = _select_anchors(evidence, close_period="2026-06", limit=6)
    keys = [k for k, _ in anchors]
    assert "period_matrix.ARR.cm.actual" in keys
    assert not any("2026-01" in k or "2026-02" in k for k in keys)


def test_repair_rendered_deck_money_whole_million_rounding() -> None:
    from app.services.reporting.export.deck_post_render_verify import (
        repair_rendered_deck_money,
        verify_rendered_deck,
    )

    raw = _mini_pptx("$80M", "$15M", "$2K", "$19M")
    evidence = {
        "deck.monthly_trends.ending_arr_m[2]": Decimal("79505000"),
        "deck.pl_detail.qtd.gross_profit.actual": Decimal("14870000"),
        "deck.gtm_performance.channels[1].spend_budget_raw": Decimal("2250"),
        "deck.cash_liquidity.current_month.cash_headroom_vs_floor": Decimal("20000000"),
        "period_matrix.ARR.cm.actual": Decimal("86100000"),
    }
    repaired, repairs = repair_rendered_deck_money(raw, evidence=evidence)
    assert repairs
    assert any("80M" in r and "79.51" in r for r in repairs)
    result = verify_rendered_deck(repaired, evidence=evidence, payload={"close_period": "2026-06"})
    assert result.numeric.ok, result.summary()


def test_pptx_script_repairs_rounded_kpi_instead_of_emdash() -> None:
    from app.services.commentary.claim_verify import apply_fail_closed_claims_to_pptx_script

    script = 'slide.addText("$80M", {x:1}); slide.addText("$2K", {x:2});'
    evidence = {
        "ending_arr": Decimal("79505000"),
        "spend_budget": Decimal("2250"),
    }
    out, result = apply_fail_closed_claims_to_pptx_script(script, evidence)
    assert "$79.51M" in out
    assert "$2.3K" in out
    assert out.count("—") == 0


def test_two_dp_millions_ties_sor_within_display_band() -> None:
    from app.services.commentary.claim_verify import verify_text_against_evidence

    result = verify_text_against_evidence(
        "$79.51M",
        {"ending_arr": Decimal("79505000")},
    )
    assert result.ok
    # Whole millions still fail — must be repaired, not waved through.
    whole = verify_text_against_evidence(
        "$80M",
        {"ending_arr": Decimal("79505000")},
    )
    assert not whole.ok


def test_post_render_narrative_filler_flagged() -> None:
    raw = _mini_pptx("ARR saw a significant expansion versus budget this quarter.")
    result = verify_rendered_deck(
        raw,
        evidence={"period_matrix.ARR.cm.actual": Decimal("86100000")},
        payload={"close_period": "2026-06"},
    )
    assert not result.narrative.ok
    assert any(i.kind == "filler" for i in result.narrative.issues)
