"""Key Takeaways stacks must not walk off the bottom of the slide.

The June-2026 MD&A deck placed its third ARR bullet at 7.45" on a 7.5" slide — entirely
invisible, with the bullet above it printed through the footer. Boxes were declared 0.55"
tall for roughly 0.20" of text, and three of them on that pitch did not fit.
"""

from __future__ import annotations

import io

import pytest
from pptx import Presentation
from pptx.util import Emu, Inches, Pt

from app.services.reporting.export.deck_layout_repair import SAFE_BOTTOM_IN, repair_deck_layout

EMU_PER_INCH = 914400
BULLET = (
    "• Expansion momentum remains strong at $869.1K, only $43.5K below budget. "
    "CS-led enterprise plays targeting $0.5-1.0M ARR/quarter are achievable at "
    "near-zero incremental CAC; recommend funding for H2 execution."
)


def _deck(tops: list[float], *, box_h: float = 0.55, with_footer: bool = True) -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    label = slide.shapes.add_textbox(Inches(0.35), Inches(tops[0] - 0.35), Inches(12.63), Inches(0.25))
    label.text_frame.text = "KEY TAKEAWAYS"
    label.text_frame.paragraphs[0].runs[0].font.size = Pt(8.5)

    for top in tops:
        box = slide.shapes.add_textbox(Inches(0.35), Inches(top), Inches(12.63), Inches(box_h))
        box.text_frame.text = BULLET
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(8.5)

    if with_footer:
        footer = slide.shapes.add_textbox(Inches(0.35), Inches(7.05), Inches(12.63), Inches(0.2))
        footer.text_frame.text = "SMPL · Board Operating Review · Q2 2026 · CONFIDENTIAL  3/11"
        footer.text_frame.paragraphs[0].runs[0].font.size = Pt(7.0)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _bottoms(data: bytes) -> list[float]:
    prs = Presentation(io.BytesIO(data))
    return [
        (sh.top + sh.height) / EMU_PER_INCH
        for s in prs.slides
        for sh in s.shapes
        if sh.top is not None and sh.height is not None
    ]


def _texts(data: bytes) -> list[str]:
    prs = Presentation(io.BytesIO(data))
    return sorted(
        sh.text_frame.text.strip()
        for s in prs.slides
        for sh in s.shapes
        if sh.has_text_frame and sh.text_frame.text.strip()
    )


def test_overflowing_stack_is_pulled_back_on_slide():
    raw = _deck([6.25, 6.85, 7.45])
    assert max(_bottoms(raw)) > 7.5, "fixture should reproduce the off-slide bullet"

    fixed, report = repair_deck_layout(raw)

    assert report.changed
    assert not report.unfixable
    assert max(_bottoms(fixed)) <= 7.5 + 0.01
    body = [b for b in _bottoms(fixed) if b < 7.05]
    assert max(body) <= SAFE_BOTTOM_IN + 0.01, "body must clear the footer"


def test_repair_preserves_every_word():
    raw = _deck([6.25, 6.85, 7.45])
    fixed, _ = repair_deck_layout(raw)
    assert _texts(fixed) == _texts(raw)


def test_repair_leaves_a_deck_that_already_fits_alone():
    raw = _deck([4.20, 4.80, 5.40])
    fixed, report = repair_deck_layout(raw)
    assert not report.changed
    assert fixed is raw


def test_footer_itself_is_not_treated_as_overflow():
    raw = _deck([4.20, 4.80, 5.40])
    _, report = repair_deck_layout(raw)
    assert not any("CONFIDENTIAL" in m for m in report.moved + report.unfixable)


def test_impossible_block_is_reported_not_silently_overlapped():
    # Ten bullets have nowhere to go below 6.5"; the pass must say so rather than
    # stack them on top of each other.
    raw = _deck([6.5 + i * 0.05 for i in range(10)])
    fixed, report = repair_deck_layout(raw)
    assert report.unfixable
    assert not report.changed
    assert fixed is raw


def test_corrupt_bytes_are_returned_untouched():
    junk = b"not a pptx"
    out, report = repair_deck_layout(junk)
    assert out is junk
    assert not report.changed


@pytest.mark.parametrize("count", [2, 3, 4])
def test_stacks_of_several_lengths_end_up_inside_the_slide(count):
    raw = _deck([6.25 + i * 0.6 for i in range(count)])
    fixed, report = repair_deck_layout(raw)
    if report.changed:
        assert max(_bottoms(fixed)) <= 7.5 + 0.01
