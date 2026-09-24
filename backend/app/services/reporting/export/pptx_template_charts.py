"""Phase 3 — inject live chart series + KPI ink into the board PPTX template.

Template-first export already rolls period labels and commentary. Without this
module, native charts and scorecard cells keep gold-deck numbers while bullets
cite the live close — the dual-ink High from the integrity review.
"""

from __future__ import annotations

import calendar
import logging
from decimal import Decimal
from typing import Any, Iterable

from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.services.reporting.export.board_chart_service import _wf
from app.services.reporting.export.board_metrics_snapshot import build_metrics_snapshot
from app.services.reporting.export.board_period_ytd import rollup_cash, rollup_ebitda, rollup_ending_arr, rollup_revenue
from app.services.reporting.export.board_platform_metrics import (
    build_arr_bridge_block,
    build_cash_liquidity_block,
    build_fy_outlook_block,
    ending_arr_at_period,
    gross_margin_pct,
)
from app.services.reporting.export.board_slide_commentary_payload import (
    fmt_deck_money,
    fmt_deck_pct,
    fmt_deck_var_pct,
)
from app.services.reporting.export.is_line_resolver import ending_cash_at
from app.services.reporting.export.period_views import ytd_periods
from app.services.reporting.export.schemas import ReportingBundle
from app.services.reporting.period_utils import prior_period, to_period

logger = logging.getLogger(__name__)

ZERO = Decimal("0")
_M = Decimal("1000000")

# Slide-key → which native chart archetype to refresh.
_CHART_BY_SLIDE_KEY: dict[str, str] = {
    "arr_waterfall": "arr_net_new",
    "gaap_revenue": "gross_profit",
    "cash_forecast": "cash_ending",
}

_SCORECARD_ROW_ALIASES: dict[str, tuple[str, ...]] = {
    "revenue": ("revenue (mrr)", "revenue"),
    "subscription_rev": ("subscription rev", "subscription revenue"),
    "gross_margin": ("gross margin",),
    "ebitda": ("adj. ebitda", "ebitda"),
    "ending_arr": ("arr (ending)", "ending arr"),
    "net_new_arr": ("net new arr",),
    "cash": ("cash (eop)", "ending cash", "cash"),
    "nrr": ("nrr (trailing)", "nrr"),
}


def _to_millions(value: Decimal | float | int | None) -> float:
    if value is None:
        return 0.0
    return float(Decimal(str(value)) / _M)


def _fmt_money(value: Decimal | None) -> str:
    """Match gold-deck accounting negatives: ($1.30M) not -$1.30M."""
    if value is None:
        return "n/a"
    raw = fmt_deck_money(value)
    if raw.startswith("-$"):
        return f"({raw[1:]})"
    return raw


def _fmt_var_money(actual: Decimal, compare: Decimal, *, label: str = "vs budget") -> str:
    delta = actual - compare
    sign = "+" if delta >= 0 else ""
    return f"{sign}{_fmt_money(delta).lstrip('+')} {label}"


def _fmt_var_pct(actual: Decimal, compare: Decimal) -> str:
    if compare == 0:
        return "—"
    return fmt_deck_var_pct(actual, compare)


def _month_abbr(period: str) -> str:
    return calendar.month_abbr[int(period[5:7])]


def _fs_line(bundle: ReportingBundle, period: str, needle: str, scenario: str) -> Decimal:
    fs = bundle.comparison_financial_statements or bundle.financial_statements
    if not fs:
        return ZERO
    n = needle.lower()
    total = ZERO
    for row in fs.income_statement.rows:
        p = to_period(str(row.period)[:7])
        if p != period or row.scenario != scenario:
            continue
        if n in row.line_item.lower() and "deferred" not in row.line_item.lower():
            total += row.amount
    return total


def _net_new_arr(bundle: ReportingBundle, period: str, scenario: str) -> Decimal:
    direct = _wf(bundle, "arr", "new_arr", period, scenario) or _wf(
        bundle, "arr", "net_new_arr", period, scenario
    )
    if direct:
        return direct
    nb = _wf(bundle, "arr", "new_business", period, scenario)
    exp = _wf(bundle, "arr", "expansion_arr", period, scenario) or _wf(
        bundle, "arr", "expansion", period, scenario
    )
    churn = abs(
        _wf(bundle, "arr", "churn_arr", period, scenario)
        or _wf(bundle, "arr", "churn", period, scenario)
    )
    cont = abs(
        _wf(bundle, "arr", "contraction_arr", period, scenario)
        or _wf(bundle, "arr", "contraction", period, scenario)
    )
    return nb + exp - churn - cont


def build_template_chart_series(
    bundle: ReportingBundle,
    archetype: str,
) -> tuple[list[str], dict[str, list[float]]] | None:
    """Return (categories, series) in $M units matching gold-deck chart scales."""
    as_of = to_period(bundle.as_of_period)
    periods = [to_period(p) for p in ytd_periods(as_of)]
    if not periods:
        return None
    cats = [_month_abbr(p) for p in periods]

    if archetype == "arr_net_new":
        act = [_to_millions(_net_new_arr(bundle, p, "Actual")) for p in periods]
        bud = [_to_millions(_net_new_arr(bundle, p, "Budget")) for p in periods]
        if not any(act) and not any(bud):
            return None
        return cats, {"Net New ARR (Act)": act, "Net New ARR (Bud)": bud}

    if archetype == "gross_profit":
        vals = [_to_millions(_fs_line(bundle, p, "gross profit", "Actual")) for p in periods]
        if not any(vals):
            # Fall back to revenue − COGS when Gross Profit line is absent.
            vals = []
            for p in periods:
                rev = _fs_line(bundle, p, "revenue", "Actual")
                cogs = _fs_line(bundle, p, "cost of", "Actual") or _fs_line(
                    bundle, p, "cogs", "Actual"
                )
                vals.append(_to_millions(rev - cogs))
        if not any(vals):
            return None
        return cats, {"Gross Profit": vals}

    if archetype == "cash_ending":
        act = [_to_millions(ending_cash_at(bundle, p, "Actual")) for p in periods]
        bud = [_to_millions(ending_cash_at(bundle, p, "Budget")) for p in periods]
        if not any(act) and not any(bud):
            return None
        return cats, {"Ending Cash (Actual)": act, "Ending Cash (Budget)": bud}

    return None


def _replace_chart_data(chart, categories: list[str], series: dict[str, list[float]]) -> bool:
    from pptx.chart.data import CategoryChartData

    data = CategoryChartData()
    data.categories = categories
    for name, values in series.items():
        padded = list(values) + [0.0] * (len(categories) - len(values))
        data.add_series(str(name), padded[: len(categories)])
    chart.replace_data(data)
    return True


def _iter_slide_charts(slide) -> Iterable[Any]:
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.CHART:
            yield shape.chart
        elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for child in shape.shapes:
                if child.shape_type == MSO_SHAPE_TYPE.CHART:
                    yield child.chart


def apply_template_charts(
    prs,
    bundle: ReportingBundle,
    key_to_idx: dict[str, int],
) -> int:
    """Replace native chart series on mapped template slides with live SoT."""
    changed = 0
    for slide_key, archetype in _CHART_BY_SLIDE_KEY.items():
        idx = key_to_idx.get(slide_key)
        if not idx:
            continue
        series_payload = build_template_chart_series(bundle, archetype)
        if not series_payload:
            logger.info("Board PPTX chart skip %s: no live series", slide_key)
            continue
        cats, series = series_payload
        slide = prs.slides[idx - 1]
        charts = list(_iter_slide_charts(slide))
        if not charts:
            logger.info("Board PPTX chart skip %s: no chart shape on slide %d", slide_key, idx)
            continue
        try:
            _replace_chart_data(charts[0], cats, series)
            changed += 1
            logger.info(
                "Board PPTX chart refreshed %s slide=%d cats=%d series=%s",
                slide_key,
                idx,
                len(cats),
                list(series.keys()),
            )
        except Exception as exc:
            logger.warning("Board PPTX chart replace failed %s: %s", slide_key, exc)
    return changed


def _shape_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    try:
        return (shape.text_frame.text or "").strip()
    except Exception:
        return ""


def _set_text_preserve_runs(shape, text: str) -> bool:
    if not getattr(shape, "has_text_frame", False):
        return False
    try:
        tf = shape.text_frame
        if not tf.paragraphs:
            return False
        p0 = tf.paragraphs[0]
        if p0.runs:
            p0.runs[0].text = text
            for run in p0.runs[1:]:
                run.text = ""
        else:
            p0.text = text
        for para in tf.paragraphs[1:]:
            for run in para.runs:
                run.text = ""
            try:
                para.text = ""
            except Exception:
                pass
        return True
    except Exception:
        return False


def _row_shapes(slide, *, min_top: int = 0, max_top: int = 10**9) -> dict[int, list[Any]]:
    from collections import defaultdict

    rows: dict[int, list[Any]] = defaultdict(list)
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        top = int(getattr(shape, "top", 0) or 0)
        if top < min_top or top > max_top:
            continue
        if not _shape_text(shape):
            continue
        key = round(top / 80000) * 80000
        rows[key].append(shape)
    for key in rows:
        rows[key].sort(key=lambda s: int(getattr(s, "left", 0) or 0))
    return rows


def _match_scorecard_row(label: str) -> str | None:
    lowered = label.strip().lower()
    for key, aliases in _SCORECARD_ROW_ALIASES.items():
        if lowered in aliases or any(a == lowered for a in aliases):
            return key
    return None


def _scorecard_values(bundle: ReportingBundle) -> dict[str, tuple[str, str, str, str, str]]:
    """Row key → (actual, budget, vs_bud%, prior, vs_prior%)."""
    as_of = to_period(bundle.as_of_period)
    prior = prior_period(as_of)
    m = build_metrics_snapshot(bundle)
    arr = rollup_ending_arr(bundle)
    rev = rollup_revenue(bundle)
    ebitda = rollup_ebitda(bundle)
    cash = rollup_cash(bundle)

    def money_row(act: Decimal, bud: Decimal, prior_act: Decimal) -> tuple[str, str, str, str, str]:
        return (
            _fmt_money(act),
            _fmt_money(bud),
            _fmt_var_pct(act, bud),
            _fmt_money(prior_act),
            _fmt_var_pct(act, prior_act) if prior_act else "—",
        )

    gm_a = gross_margin_pct(bundle, as_of, "Actual")
    gm_b = gross_margin_pct(bundle, as_of, "Budget")
    gm_p = gross_margin_pct(bundle, prior, "Actual")
    if gm_a is not None and gm_b is not None:
        gm_var = f"{(float(gm_a) - float(gm_b)):+.1f}pp"
    else:
        gm_var = "—"

    nn = m.net_new_arr
    nn_b = m.new_arr_budget
    nn_p = _net_new_arr(bundle, prior, "Actual")
    nrr_disp = fmt_deck_pct(m.nrr, as_percent=False) if m.nrr is not None else "n/a"

    return {
        "revenue": money_row(
            rev.current_month, rev.budget_cm, _fs_line(bundle, prior, "revenue", "Actual")
        ),
        "subscription_rev": money_row(
            rev.current_month, rev.budget_cm, _fs_line(bundle, prior, "revenue", "Actual")
        ),
        "gross_margin": (
            fmt_deck_pct(gm_a, as_percent=False) if gm_a is not None else "n/a",
            fmt_deck_pct(gm_b, as_percent=False) if gm_b is not None else "n/a",
            gm_var,
            fmt_deck_pct(gm_p, as_percent=False) if gm_p is not None else "n/a",
            "—",
        ),
        "ebitda": money_row(
            ebitda.current_month, ebitda.budget_cm, _fs_line(bundle, prior, "ebitda", "Actual")
        ),
        "ending_arr": money_row(
            arr.current_month, arr.budget_cm, ending_arr_at_period(bundle, prior, "Actual")
        ),
        "net_new_arr": money_row(nn, nn_b, nn_p),
        "cash": money_row(
            cash.current_month, cash.budget_cm, ending_cash_at(bundle, prior, "Actual")
        ),
        "nrr": (nrr_disp, "n/a", "—", "n/a", "—"),
    }


def apply_executive_scorecard(prs, bundle: ReportingBundle, key_to_idx: dict[str, int]) -> int:
    idx = key_to_idx.get("executive_summary")
    if not idx:
        return 0
    slide = prs.slides[idx - 1]
    values = _scorecard_values(bundle)
    changed = 0
    for _top, shapes in _row_shapes(slide, min_top=1_200_000, max_top=3_900_000).items():
        if len(shapes) < 2:
            continue
        row_key = _match_scorecard_row(_shape_text(shapes[0]))
        if not row_key or row_key not in values:
            continue
        cells = values[row_key]
        # Label + up to 5 value cells to the right.
        targets = shapes[1:6]
        for shape, text in zip(targets, cells):
            if _set_text_preserve_runs(shape, text):
                changed += 1
    return changed


def _column_groups(slide, *, max_top: int) -> list[list[Any]]:
    """Group text shapes into left-to-right KPI cards by left edge."""
    from collections import defaultdict

    cols: dict[int, list[Any]] = defaultdict(list)
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        top = int(getattr(shape, "top", 0) or 0)
        if top > max_top or top < 900_000:
            continue
        text = _shape_text(shape)
        if not text:
            continue
        key = round(int(getattr(shape, "left", 0) or 0) / 400_000) * 400_000
        cols[key].append(shape)
    groups: list[list[Any]] = []
    for key in sorted(cols):
        group = sorted(cols[key], key=lambda s: int(getattr(s, "top", 0) or 0))
        if len(group) >= 2:
            groups.append(group)
    return groups


def _update_kpi_card(group: list[Any], *, value: str, subtext: str | None = None) -> int:
    """Update value (2nd text) and optional subtext (3rd) under a label."""
    changed = 0
    if len(group) >= 2 and _set_text_preserve_runs(group[1], value):
        changed += 1
    if subtext is not None and len(group) >= 3 and _set_text_preserve_runs(group[2], subtext):
        changed += 1
    return changed


def apply_slide_kpi_cards(prs, bundle: ReportingBundle, key_to_idx: dict[str, int]) -> int:
    """Refresh top KPI cards on ARR / P&L / Cash slides from live SoT."""
    as_of = to_period(bundle.as_of_period)
    m = build_metrics_snapshot(bundle)
    arr_block = build_arr_bridge_block(bundle)
    cash_block = build_cash_liquidity_block(bundle)
    fy = build_fy_outlook_block(bundle)
    rev = rollup_revenue(bundle)
    ebitda = rollup_ebitda(bundle)
    gm = gross_margin_pct(bundle, as_of, "Actual")
    changed = 0

    # --- ARR slide ---
    idx = key_to_idx.get("arr_waterfall")
    if idx:
        slide = prs.slides[idx - 1]
        for group in _column_groups(slide, max_top=1_900_000):
            label = _shape_text(group[0]).lower()
            if "fy arr" in label or ("forecast" in label and "arr" in label):
                arr_eoy = fy.get("arr_eoy") or {}
                value = (
                    arr_eoy.get("actual_or_outlook")
                    or arr_block.get("arr_eop")
                    or _fmt_money(m.ending_arr)
                )
                bud = ending_arr_at_period(bundle, as_of, "Budget")
                eoy_raw = ending_arr_at_period(bundle, as_of[:4] + "-12", "Forecast") or m.ending_arr
                sub = _fmt_var_money(eoy_raw, bud, label="vs budget") if bud else None
                changed += _update_kpi_card(group, value=str(value), subtext=sub)
            elif "net new" in label:
                changed += _update_kpi_card(
                    group,
                    value=arr_block.get("net_new") or _fmt_money(m.net_new_arr),
                    subtext=_fmt_var_money(m.net_new_arr, m.new_arr_budget, label="vs plan"),
                )
            elif "expansion" in label:
                changed += _update_kpi_card(
                    group,
                    value=arr_block.get("expansion") or "n/a",
                    subtext=None,
                )
            elif "churn" in label:
                changed += _update_kpi_card(
                    group,
                    value=arr_block.get("churn") or "n/a",
                    subtext=None,
                )

    # --- P&L slide ---
    idx = key_to_idx.get("gaap_revenue")
    if idx:
        slide = prs.slides[idx - 1]
        for group in _column_groups(slide, max_top=1_900_000):
            label = _shape_text(group[0]).lower()
            if "revenue" in label and "ytd" not in label:
                changed += _update_kpi_card(
                    group,
                    value=_fmt_money(rev.current_month),
                    subtext=_fmt_var_money(rev.current_month, rev.budget_cm, label="vs bud"),
                )
            elif "ytd" in label:
                changed += _update_kpi_card(
                    group,
                    value=_fmt_money(rev.ytd),
                    subtext=_fmt_var_money(rev.ytd, rev.budget_ytd, label="vs bud"),
                )
            elif "margin" in label or "gm" in label:
                changed += _update_kpi_card(
                    group,
                    value=fmt_deck_pct(gm, as_percent=False) if gm is not None else "n/a",
                    subtext=None,
                )
            elif "ebitda" in label:
                changed += _update_kpi_card(
                    group,
                    value=_fmt_money(ebitda.current_month),
                    subtext=_fmt_var_money(ebitda.current_month, ebitda.budget_cm, label="vs bud"),
                )

    # --- Cash slide ---
    idx = key_to_idx.get("cash_forecast")
    if idx:
        cm = cash_block.get("current_month") or {}
        slide = prs.slides[idx - 1]
        for group in _column_groups(slide, max_top=1_900_000):
            label = _shape_text(group[0]).lower()
            if "cash" in label and ("eop" in label or "ending" in label or label.strip() == "cash"):
                act = cm.get("cash_eop_actual") or _fmt_money(rollup_cash(bundle).current_month)
                bud = cm.get("cash_eop_budget")
                sub = None
                if bud and bud not in {"—", "n/a"}:
                    sub = f"{act} vs budget"  # overwritten below with delta
                    try:
                        sub = _fmt_var_money(
                            ending_cash_at(bundle, as_of, "Actual"),
                            ending_cash_at(bundle, as_of, "Budget"),
                            label="vs budget",
                        )
                    except Exception:
                        sub = None
                changed += _update_kpi_card(group, value=str(act), subtext=sub)
            elif "headroom" in label or "floor" in label:
                changed += _update_kpi_card(
                    group,
                    value=str(cm.get("cash_headroom_vs_floor") or "n/a"),
                    subtext=f"Floor: {cm.get('cash_floor') or '$10.0M'}",
                )
            elif "cfo" in label or "operating" in label:
                changed += _update_kpi_card(
                    group,
                    value=str(cm.get("cfo_actual") or "n/a"),
                    subtext=None,
                )

    return changed


def apply_template_visuals(
    prs,
    bundle: ReportingBundle,
    key_to_idx: dict[str, int] | None = None,
) -> dict[str, int]:
    """Apply charts + KPI ink. Returns counts by surface."""
    if key_to_idx is None:
        from app.services.reporting.export.pptx_template_export import map_template_slides_to_slide_keys

        key_to_idx = map_template_slides_to_slide_keys(prs)

    charts = apply_template_charts(prs, bundle, key_to_idx)
    scorecard = apply_executive_scorecard(prs, bundle, key_to_idx)
    kpis = apply_slide_kpi_cards(prs, bundle, key_to_idx)
    logger.info(
        "Board PPTX Phase 3 visuals: charts=%d scorecard_cells=%d kpi_cells=%d",
        charts,
        scorecard,
        kpis,
    )
    return {"charts": charts, "scorecard": scorecard, "kpis": kpis}
