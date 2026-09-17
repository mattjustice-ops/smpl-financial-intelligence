"""Post-render MD&A deck fidelity — compare PPTX cells to the evidence payload.

Pre-render claim_verify only sees PptxGenJS string literals. After Node renders
the file, charts, merged shapes, and postprocess reinjection can still diverge.
This module extracts visible slide text from the finished PPTX and:

1. Re-runs numeric claim verify on short metric cells (same TOL_ACTUALS).
2. Checks that high-priority payload anchors (ARR / revenue / cash) appear
   somewhere in the deck within tolerance.
3. Runs non-numeric narrative checks on longer prose shapes.

Policy for Prompt 5: soft-warn + export (never block the customer download).
Callers log ``result.summary()``; optional hard-block via ``fail_closed=True``.
"""

from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Any, Mapping, Sequence

from app.services.commentary.claim_verify import (
    TOL_ACTUALS,
    ClaimCheck,
    CommentaryIntegrityError,
    NumericClaim,
    VerificationResult,
    _best_match,
    _pptx_is_narrative_literal,
    _pptx_material_claims,
    _to_decimal,
    evidence_values_from_package,
)
from app.services.commentary.narrative_verify import (
    NarrativeIssue,
    NarrativeVerificationResult,
    evidence_close_period,
    verify_narrative_text,
)

logger = logging.getLogger(__name__)

_METRIC_CELL_RE = re.compile(
    r"^[\s\$\(\+\-–—]*"
    r"(?:"
    r"\d{1,3}(?:,\d{3})+(?:\.\d+)?[MmKkBb]?|"
    r"\d+(?:\.\d+)?[MmKkBb]?|"
    r"\d+(?:\.\d+)?%|"
    r"\d+(?:\.\d+)?[xX]|"
    r"[Nn]/?[Aa]|—"
    r")"
    r"[\s\)\%xXMmKkBb]*$"
)

# Prefer these payload paths as "must appear somewhere" anchors when present.
_ANCHOR_KEY_HINTS: tuple[str, ...] = (
    "ending_arr",
    "arr.cm.actual",
    "period_matrix.ARR.cm.actual",
    "revenue",
    "period_matrix.Revenue.cm.actual",
    "ending_cash",
    "period_matrix.Cash.cm.actual",
    "net_new",
)


@dataclass
class DeckPostRenderResult:
    numeric: VerificationResult = field(default_factory=VerificationResult)
    anchors: VerificationResult = field(default_factory=VerificationResult)
    narrative: NarrativeVerificationResult = field(default_factory=NarrativeVerificationResult)
    slide_count: int = 0
    text_shape_count: int = 0

    @property
    def ok(self) -> bool:
        return self.numeric.ok and self.anchors.ok and self.narrative.ok

    def summary(self, *, max_failures: int = 8) -> str:
        parts: list[str] = []
        if not self.numeric.ok:
            parts.append(f"cells: {self.numeric.summary(max_failures=max_failures)}")
        if not self.anchors.ok:
            parts.append(f"anchors: {self.anchors.summary(max_failures=max_failures)}")
        if not self.narrative.ok:
            parts.append(f"narrative: {self.narrative.summary(max_issues=max_failures)}")
        if not parts:
            return (
                f"post-render ok ({self.slide_count} slides, "
                f"{self.text_shape_count} text shapes)"
            )
        return " | ".join(parts)


def _iter_shapes(shapes: Any):
    for shape in shapes:
        yield shape
        if hasattr(shape, "shapes"):
            yield from _iter_shapes(shape.shapes)


def extract_pptx_texts(pptx_bytes: bytes) -> tuple[int, list[str]]:
    """Return (slide_count, non-empty shape texts) from a PPTX byte blob."""
    from pptx import Presentation

    prs = Presentation(io.BytesIO(pptx_bytes))
    texts: list[str] = []
    for slide in prs.slides:
        for shape in _iter_shapes(slide.shapes):
            if not hasattr(shape, "text_frame") or shape.text_frame is None:
                continue
            parts = [p.text for p in shape.text_frame.paragraphs if p.text and p.text.strip()]
            blob = "\n".join(parts).strip()
            if blob:
                texts.append(blob)
    return len(prs.slides), texts


def _evidence_map(evidence: Mapping[str, Decimal] | Mapping[str, Any]) -> dict[str, Decimal]:
    if hasattr(evidence, "get") and "values_decimal" in evidence:
        return evidence_values_from_package(evidence)  # type: ignore[arg-type]
    out: dict[str, Decimal] = {}
    for k, v in evidence.items():
        num = v if isinstance(v, Decimal) else _to_decimal(v)
        if num is not None:
            out[str(k)] = num
    return out


def _flatten_payload_evidence(payload: Mapping[str, Any] | None) -> dict[str, Decimal]:
    if not payload:
        return {}
    from app.services.reporting.export.board_platform_metrics import (
        build_evidence_package_from_deck_payload,
        evidence_values_from_deck_payload,
    )

    try:
        package = build_evidence_package_from_deck_payload(dict(payload))
        values = evidence_values_from_package(package) or evidence_values_from_deck_payload(
            dict(payload)
        )
        return dict(values or {})
    except Exception as exc:  # pragma: no cover — defensive
        logger.warning("post-render: could not build evidence from payload: %s", exc)
        return {}


def _anchor_score(key: str, close_period: str | None) -> int:
    """Score anchors that appear as on-slide KPI text, not chart-only monthly series."""
    kl = key.lower()
    # Monthly trend story keys (Jan–May ARR etc.) live in chart series only
    # (showValue:false) — they never appear as $ text and create false anchor misses.
    if close_period and "ending_arr" in kl and re.search(r"\d{4}-\d{2}", kl):
        periods = re.findall(r"\d{4}-\d{2}", kl)
        if periods and all(p != close_period for p in periods):
            return 0
    if "monthly_trends" in kl and close_period:
        periods = re.findall(r"\d{4}-\d{2}", kl)
        if periods and all(p != close_period for p in periods):
            return 0
    score = 0
    for i, hint in enumerate(_ANCHOR_KEY_HINTS):
        if hint.lower() in kl:
            score = 100 - i
            break
    if score <= 0:
        return 0
    # Prefer period_matrix / executive / CM close KPIs that Claude copies into cells.
    if "period_matrix" in kl or ".cm." in kl or "executive" in kl:
        score += 40
    if close_period and close_period in kl:
        score += 50
    return score


def _select_anchors(
    evidence: Mapping[str, Decimal],
    *,
    limit: int = 6,
    close_period: str | None = None,
) -> list[tuple[str, Decimal]]:
    """Pick a small set of high-priority evidence values that should appear in the deck."""
    ranked: list[tuple[int, str, Decimal]] = []
    for key, val in evidence.items():
        score = _anchor_score(key, close_period)
        if score <= 0:
            continue
        # Prefer absolute dollars in board display range.
        if abs(val) < Decimal("1000"):
            continue
        ranked.append((score, key, val))
    ranked.sort(key=lambda t: (-t[0], t[1]))
    seen: set[Decimal] = set()
    out: list[tuple[str, Decimal]] = []
    for _, key, val in ranked:
        # Dedupe near-equal amounts so one ARR figure satisfies one anchor.
        rounded = val.quantize(Decimal("1"))
        if rounded in seen:
            continue
        seen.add(rounded)
        out.append((key, val))
        if len(out) >= limit:
            break
    return out


def _rounding_repair_band(claim: NumericClaim) -> Decimal | None:
    """Max |claim − SoR| we will auto-correct as display rounding (not invention)."""
    if claim.kind != "money":
        return None
    stated = claim.stated or ""
    # Whole millions/thousands ($80M, $2K) are the common Claude/adapt defect.
    if re.search(r"\dM\b", stated, re.I) and not re.search(r"\.\d+M\b", stated, re.I):
        return Decimal("1500000")  # within $1.5M of SoR
    if re.search(r"\dK\b", stated, re.I) and not re.search(r"\.\d+K\b", stated, re.I):
        return Decimal("1000")
    if re.search(r"\.\d{1,2}M\b", stated, re.I):
        return Decimal("10000")  # 2dp M half-ulp neighborhood
    if re.search(r"\.\d{1,2}K\b", stated, re.I):
        return Decimal("50")
    return Decimal("1")


def _canonical_money_display(value: Decimal) -> str:
    from app.services.reporting.export.board_slide_commentary_payload import fmt_deck_money

    return fmt_deck_money(value)


def repair_rendered_deck_money(
    pptx_bytes: bytes,
    *,
    payload: Mapping[str, Any] | None = None,
    evidence: Mapping[str, Decimal] | Mapping[str, Any] | None = None,
) -> tuple[bytes, list[str]]:
    """Rewrite rounded KPI cells in the finished PPTX to SoR-canonical display strings.

    Analysts tie prior periods to the system of record; whole-million / whole-thousand
    Claude rounding ($80M vs $79.505M) must not ship. Only near-miss display rounding
    inside ``_rounding_repair_band`` is corrected — distant inventions stay for verify.
    """
    from pptx import Presentation

    evidence_map = _evidence_map(evidence) if evidence is not None else _flatten_payload_evidence(
        payload
    )
    if not evidence_map:
        return pptx_bytes, []

    prs = Presentation(io.BytesIO(pptx_bytes))
    repairs: list[str] = []

    for slide in prs.slides:
        for shape in _iter_shapes(slide.shapes):
            if not hasattr(shape, "text_frame") or shape.text_frame is None:
                continue
            for para in shape.text_frame.paragraphs:
                raw = para.text or ""
                if not raw.strip():
                    continue
                # Repair short metric cells and short multi-token KPI lines; leave long narrative.
                if _pptx_is_narrative_literal(raw) and not _METRIC_CELL_RE.match(raw.strip()):
                    continue
                new_text = raw
                for claim in _pptx_material_claims(raw):
                    if claim.kind != "money":
                        continue
                    check = _best_match(claim, evidence_map)
                    if check.status == "pass":
                        continue
                    nearest = check.matched_value
                    band = _rounding_repair_band(claim)
                    if nearest is None or band is None or check.diff is None:
                        continue
                    if check.diff > band:
                        continue
                    canon = _canonical_money_display(nearest)
                    if canon in {"n/a", "—"} or canon == claim.stated:
                        continue
                    if claim.stated not in new_text:
                        continue
                    new_text = new_text.replace(claim.stated, canon, 1)
                    repairs.append(f"{claim.stated}→{canon}")
                if new_text != raw:
                    para.text = new_text

    if not repairs:
        return pptx_bytes, []
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue(), repairs


def _money_values_in_texts(texts: Sequence[str]) -> list[Decimal]:
    found: list[Decimal] = []
    for text in texts:
        for claim in _pptx_material_claims(text):
            if claim.kind == "money":
                found.append(claim.value)
    return found


def _anchor_present(target: Decimal, money_values: Sequence[Decimal]) -> bool:
    for v in money_values:
        if abs(v - target) <= TOL_ACTUALS:
            return True
        # Display often rounds to $0.1M — allow 50k band for large anchors.
        if abs(target) >= Decimal("1000000") and abs(v - target) <= Decimal("50000"):
            return True
    return False


def verify_rendered_deck(
    pptx_bytes: bytes,
    *,
    payload: Mapping[str, Any] | None = None,
    evidence: Mapping[str, Decimal] | Mapping[str, Any] | None = None,
    fail_closed: bool = False,
) -> DeckPostRenderResult:
    """Compare finished PPTX text to evidence / payload anchors + narrative rules."""
    slide_count, texts = extract_pptx_texts(pptx_bytes)
    evidence_map = _evidence_map(evidence) if evidence is not None else _flatten_payload_evidence(
        payload
    )

    cell_checks: list[ClaimCheck] = []
    narrative_issues: list[NarrativeIssue] = []
    close_period = evidence_close_period(payload)

    for text in texts:
        stripped = text.strip()
        if _METRIC_CELL_RE.match(stripped) or not _pptx_is_narrative_literal(stripped):
            for claim in _pptx_material_claims(stripped):
                if claim.kind == "money" or claim.kind == "percent" or (
                    claim.kind == "ratio" and "x" in claim.stated.lower()
                ):
                    cell_checks.append(_best_match(claim, evidence_map))
        else:
            local = verify_narrative_text(stripped, close_period=close_period)
            narrative_issues.extend(local.issues)

    money_values = _money_values_in_texts(texts)
    anchor_checks: list[ClaimCheck] = []
    for key, val in _select_anchors(evidence_map, close_period=close_period):
        fake_claim = NumericClaim(stated=f"anchor:{key}", value=val, kind="money")
        if _anchor_present(val, money_values):
            anchor_checks.append(
                ClaimCheck(
                    claim=fake_claim,
                    status="pass",
                    matched_key=key,
                    matched_value=val,
                    diff=Decimal("0"),
                )
            )
        else:
            anchor_checks.append(
                ClaimCheck(
                    claim=fake_claim,
                    status="missing_evidence" if not money_values else "mismatch",
                    matched_key=key,
                    matched_value=None,
                    diff=None,
                )
            )

    result = DeckPostRenderResult(
        numeric=VerificationResult(checks=cell_checks),
        anchors=VerificationResult(checks=anchor_checks),
        narrative=NarrativeVerificationResult(issues=narrative_issues),
        slide_count=slide_count,
        text_shape_count=len(texts),
    )
    if fail_closed and not result.ok:
        raise CommentaryIntegrityError(
            "P15 fail-closed: post-render deck fidelity failed. " + result.summary(),
        )
    return result


def verify_rendered_deck_soft(
    pptx_bytes: bytes,
    *,
    payload: Mapping[str, Any] | None = None,
) -> tuple[bytes, DeckPostRenderResult]:
    """Prompt 5 path: repair display-rounding, then soft-warn; never raise.

    Returns ``(possibly_repaired_pptx_bytes, result)`` so the shipped file matches SoR.
    """
    repaired, repairs = repair_rendered_deck_money(pptx_bytes, payload=payload)
    if repairs:
        logger.info(
            "P15 post-render money repair (%s): %s",
            len(repairs),
            "; ".join(repairs[:12]),
        )
    result = verify_rendered_deck(repaired, payload=payload, fail_closed=False)
    if result.ok:
        logger.info("P15 post-render fidelity passed: %s", result.summary())
    else:
        logger.warning(
            "P15 post-render fidelity soft-warn (export continues): %s",
            result.summary(),
        )
    return repaired, result
