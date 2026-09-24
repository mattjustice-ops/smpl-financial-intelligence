"""Build SEO Keywords Before/After Excel from git history + current seed.

Before state is RECONSTRUCTED from git (no formal keyword list existed).
"""
from __future__ import annotations

import json
import re
import subprocess
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "marketing" / "SEO_Keywords_Before_After_2026-08.xlsx"


def git_show(spec: str) -> str:
    r = subprocess.run(
        ["git", "show", spec],
        cwd=ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    return r.stdout if r.returncode == 0 else ""


def extract_posts_from_js(text: str) -> list[dict]:
    """Pull post SEO fields from JS/JSON-ish seed source."""
    posts: list[dict] = []
    # Split on post document starts
    chunks = re.split(r"(?=_id:\s*[\"']post-|_id[\"']:\s*[\"']post-)", text)
    for chunk in chunks:
        if "post-" not in chunk[:80] and not re.search(
            r"_id[\"']?\s*:\s*[\"']post-", chunk[:120]
        ):
            # also handle _id first line after split
            if not re.match(r"_id", chunk.strip()[:20]):
                continue
        pid_m = re.search(r"_id[\"']?\s*:\s*[\"'](post-[^\"']+)[\"']", chunk)
        if not pid_m:
            continue
        # Only take until next top-level export or clear end; truncate body noise
        block = chunk[:4000]
        if "_type" in block and "post" not in block[:200]:
            continue
        if not re.search(r"_type[\"']?\s*:\s*[\"']post[\"']", block):
            # JSON posts may have _type after other fields
            if '"_type": "post"' not in block and "_type: \"post\"" not in block:
                continue

        def grab(pat: str) -> str:
            m = re.search(pat, block, re.S)
            if not m:
                return ""
            return " ".join(m.group(1).split())

        title = grab(r"[\"']?title[\"']?\s*:\s*[\"']([^\"']+)[\"']")
        slug = grab(
            r"[\"']?slug[\"']?\s*:\s*(?:\{\s*[\"']?_type[\"']?\s*:\s*[\"']slug[\"']\s*,\s*[\"']?current[\"']?\s*:\s*)?[\"']([^\"']+)[\"']"
        )
        # multiline slug.current
        if not slug:
            slug = grab(
                r"slug:\s*\{\s*_type:\s*[\"']slug[\"'],\s*current:\s*[\"']([^\"']+)[\"']"
            )
        if not slug:
            slug = grab(
                r"slug:\s*\{\s*\n\s*_type:\s*[\"']slug[\"'],\s*\n\s*current:\s*[\"']([^\"']+)[\"']"
            )
        seo_t = grab(r"[\"']?seoTitle[\"']?\s*:\s*[\"']([^\"']+)[\"']")
        seo_d = grab(r"[\"']?seoDescription[\"']?\s*:\s*[\"']([^\"']*)[\"']")
        if title and slug:
            posts.append(
                {
                    "id": pid_m.group(1),
                    "title": title,
                    "slug": slug,
                    "seoTitle": seo_t,
                    "seoDescription": seo_d,
                }
            )
    # Dedup by slug
    by_slug = {p["slug"]: p for p in posts}
    return sorted(by_slug.values(), key=lambda x: x["slug"])


def extract_glossary(text: str) -> list[dict]:
    terms = []
    # Seed glossary objects often omit _type (added at write time).
    for m in re.finditer(
        r"_id:\s*[\"'](glossary-[^\"']+)[\"'],\s*"
        r"(?:_type:\s*[\"']glossaryTerm[\"'],\s*)?"
        r"term:\s*[\"']([^\"']+)[\"'],\s*"
        r"slug:\s*[\"']([^\"']+)[\"'],\s*"
        r"shortDefinition:\s*(?:\n\s*)?[\"']([^\"']*)[\"']",
        text,
    ):
        terms.append(
            {
                "id": m.group(1),
                "term": m.group(2),
                "slug": m.group(3),
                "shortDefinition": " ".join(m.group(4).split()),
            }
        )
    by = {t["slug"]: t for t in terms}
    return sorted(by.values(), key=lambda x: x["slug"])


def infer_primary_keyword(title: str, seo_title: str = "") -> str:
    t = (seo_title or title or "").replace(" | SMPL.ai", "").strip()
    # Heuristic: drop brand suffix and take core phrase
    drops = [
        "What CFOs Should Demand From ",
        "Why ",
        "The Hidden Cost of ",
        "What is an ",
        "10 ",
    ]
    out = t
    for d in drops:
        if out.startswith(d):
            out = out[len(d) :]
            break
    return out.strip(" ?.")


def style_header(ws, row=1):
    fill = PatternFill("solid", fgColor="0F766E")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def autosize(ws, max_width=56):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = 12
        for cell in col:
            if cell.value:
                width = min(max_width, max(width, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = width


def write_rows(ws, headers, rows):
    ws.append(headers)
    style_header(ws)
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    autosize(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    # --- BEFORE: reconstructed from git ---
    before_pages = [
        {
            "surface": "Homepage / marketing default",
            "url_path": "/",
            "title": "SMPL.ai · AI operating system for SaaS finance teams",
            "meta_description": "SMPL connects pipeline, ARR, revenue, cash, headcount, and financial statements into one intelligent operating model — governed, board-ready, and built for executive decisions.",
            "primary_keyword": "AI operating system for SaaS finance",
            "source": "frontend/app/(marketing)/layout.tsx @ 2655a8c^",
        },
        {
            "surface": "Root layout default title",
            "url_path": "(global default)",
            "title": "SMPL.ai · CFO Operating Intelligence",
            "meta_description": "AI operating system for SaaS finance teams — pipeline, ARR, cash, and board-ready reporting in one model.",
            "primary_keyword": "CFO Operating Intelligence / AI OS for SaaS finance",
            "source": "frontend/app/layout.tsx + frontend/lib/site.ts @ 2655a8c^",
        },
        {
            "surface": "Book demo",
            "url_path": "/book-demo",
            "title": "Book a demo · SMPL.ai",
            "meta_description": "Share your contact details and priorities, then schedule a live SMPL demo with our team.",
            "primary_keyword": "book demo (weak / brand-only)",
            "source": "frontend/app/(marketing)/book-demo/page.tsx @ 2655a8c^",
        },
        {
            "surface": "Request quote",
            "url_path": "/request-quote",
            "title": "Request a quote · SMPL.ai",
            "meta_description": "A few quick questions about your stack and goals. SMPL will suggest the right package and coordinate a follow-up on your schedule.",
            "primary_keyword": "request quote (weak / brand-only)",
            "source": "frontend/app/(marketing)/request-quote/page.tsx @ 2655a8c^",
        },
        {
            "surface": "Pricing",
            "url_path": "/pricing",
            "title": "(inherited marketing default — no page-level metadata)",
            "meta_description": "(inherited — no unique pricing SEO)",
            "primary_keyword": "(none unique)",
            "source": "pricing/page.tsx had no Metadata export before 2655a8c",
        },
    ]

    initial_seed = git_show("8e5517f:frontend/sanity/seed/content.mjs")
    before_posts = extract_posts_from_js(initial_seed)
    # Fallback hardcode if parser misses (known from git log / transcript)
    if len(before_posts) < 3:
        before_posts = [
            {
                "id": "post-board-numbers-need-evidence",
                "title": "Why board numbers need evidence, not another dashboard",
                "slug": "board-numbers-need-evidence",
                "seoTitle": "Board Numbers Need Evidence, Not Another Dashboard | SMPL.ai",
                "seoDescription": "",
            },
            {
                "id": "post-saas-close-load-validate-lock-freeze",
                "title": "SaaS close: Load → Validate → Lock → Freeze",
                "slug": "saas-close-load-validate-lock-freeze",
                "seoTitle": "SaaS Close: Load → Validate → Lock → Freeze | SMPL.ai",
                "seoDescription": "",
            },
            {
                "id": "post-ai-commentary-finance-will-sign",
                "title": "AI commentary that finance will actually sign",
                "slug": "ai-commentary-finance-will-sign",
                "seoTitle": "AI Commentary That Finance Will Actually Sign | SMPL.ai",
                "seoDescription": "",
            },
        ]
    # Fill descriptions from git show if empty
    for p in before_posts:
        if not p.get("seoDescription"):
            # pull from known snippets in initial seed via regex on full file
            m = re.search(
                rf"slug:\s*\{{[^}}]*current:\s*[\"']{re.escape(p['slug'])}[\"'][^}}]*\}}.*?seoDescription:\s*[\"']([^\"']*)[\"']",
                initial_seed,
                re.S,
            )
            if m:
                p["seoDescription"] = " ".join(m.group(1).split())

    # Get descriptions for initial posts more carefully from git
    for slug, fallback_desc in [
        (
            "board-numbers-need-evidence",
            "Why SaaS board reporting fails without evidence trails — and how governed, traceable numbers rebuild executive trust.",
        ),
        (
            "saas-close-load-validate-lock-freeze",
            "A practical SaaS close sequence: Load → Validate → Lock → Freeze so board numbers stop moving mid-meeting.",
        ),
        (
            "ai-commentary-finance-will-sign",
            "What it takes for AI variance commentary to meet a CFO sign-off bar — evidence, drivers, and governance.",
        ),
    ]:
        for p in before_posts:
            if p["slug"] == slug and not p.get("seoDescription"):
                # try excerpt from seed
                m = re.search(
                    rf"current:\s*[\"']{slug}[\"'].{{0,400}}?excerpt:\s*[\"']([^\"']+)[\"']",
                    initial_seed,
                    re.S,
                )
                p["seoDescription"] = (
                    " ".join(m.group(1).split()) if m else fallback_desc
                )

    before_glossary = extract_glossary(initial_seed)
    if not before_glossary:
        # known set from commit 8e5517f / seed-sanity legacy IDs
        before_glossary = [
            {"slug": s, "term": t, "shortDefinition": "", "id": f"glossary-{s}"}
            for s, t in [
                ("arr", "ARR"),
                ("mrr", "MRR"),
                ("nrr", "NRR"),
                ("grr", "GRR"),
                ("deferred-revenue", "Deferred revenue"),
                ("waterfall", "Waterfall"),
                ("close", "Close"),
                ("mda", "MD&A"),
                ("freeze-pack", "Freeze pack"),
                ("combined-scenario", "Combined scenario"),
                ("bookings", "Bookings"),
                ("pipeline", "Pipeline"),
                ("churn", "Churn"),
                ("expansion", "Expansion"),
                ("contraction", "Contraction"),
                ("cac", "CAC"),
                ("ltv", "LTV"),
                ("burn-multiple", "Burn multiple"),
                ("runway", "Runway"),
                ("gaap-revenue", "GAAP revenue"),
            ]
        ]

    before_glossary_meta = {
        "surface": "Glossary index",
        "url_path": "/glossary",
        "title": "Glossary | SMPL.ai",
        "meta_description": "SaaS FP&A glossary: ARR, NRR, deferred revenue, waterfall, close, MD&A, freeze pack, and more.",
        "primary_keyword": "SaaS FP&A glossary (+ freeze pack)",
        "source": "glossary/page.tsx @ 8e5517f / pre-3577275",
    }

    # --- AFTER: current repo ---
    after_pages = [
        {
            "surface": "Homepage / SITE_TITLE",
            "url_path": "/",
            "title": "SaaS FP&A Software & Board Reporting, Built to Be Trusted | SMPL.ai",
            "meta_description": "FP&A for SaaS finance teams. Unify ARR, pipeline, cash, and financial statements into one governed model for close. Every number board-ready and traceable to its source.",
            "primary_keyword": "SaaS FP&A software; board reporting",
            "source": "frontend/lib/site.ts (current; set in 2655a8c 2026-07-21)",
        },
        {
            "surface": "Book demo",
            "url_path": "/book-demo",
            "title": "Book a Demo | SMPL.ai SaaS FP&A Platform",
            "meta_description": "See SMPL.ai live: SaaS FP&A, ARR waterfalls, close reporting, and board decks. Book a demo with our team.",
            "primary_keyword": "SaaS FP&A platform demo",
            "source": "book-demo/page.tsx (current)",
        },
        {
            "surface": "Pricing",
            "url_path": "/pricing",
            "title": "SMPL.ai Pricing | SaaS FP&A Plans",
            "meta_description": "Compare SMPL.ai plans for SaaS FP&A, ARR reporting, close packages, and board decks. Starter, Professional, and Enterprise.",
            "primary_keyword": "SaaS FP&A pricing / plans",
            "source": "pricing/page.tsx (current)",
        },
        {
            "surface": "Request quote",
            "url_path": "/request-quote",
            "title": "Request a Quote | SMPL.ai Pricing & Packaging",
            "meta_description": "Tell us your finance stack and goals. We'll recommend the right SMPL.ai package for SaaS FP&A, reporting, and board readiness.",
            "primary_keyword": "SaaS FP&A pricing quote",
            "source": "request-quote/page.tsx (current)",
        },
        {
            "surface": "Blog index",
            "url_path": "/blog",
            "title": "Blog | SMPL.ai",
            "meta_description": "Insights on SaaS FP&A, board reporting, close workflow, and AI commentary that finance teams can sign.",
            "primary_keyword": "SaaS FP&A blog / board reporting / AI commentary",
            "source": "blog/page.tsx (current)",
        },
        {
            "surface": "Glossary index",
            "url_path": "/glossary",
            "title": "Glossary | SMPL.ai",
            "meta_description": "SaaS FP&A glossary: ARR, NRR, deferred revenue, waterfall, close, MD&A, FP&A, and more.",
            "primary_keyword": "SaaS FP&A glossary",
            "source": "glossary/page.tsx (current; freeze pack removed 3577275)",
        },
        {
            "surface": "Login (noindex)",
            "url_path": "/login",
            "title": "Sign in | SMPL.ai",
            "meta_description": "Sign in to your SMPL.ai workspace.",
            "primary_keyword": "(noindex — not a ranking target)",
            "source": "login/page.tsx (current)",
        },
    ]

    # Load current posts via node for accuracy
    node_script = r"""
import { posts, glossaryTerms } from '../frontend/sanity/seed/content.mjs';
import { july29Posts } from '../frontend/sanity/seed/july29-posts.mjs';
import { july30Posts } from '../frontend/sanity/seed/july30-posts.mjs';
import { writeFileSync } from 'fs';
import { resolve, dirname } from 'path';
import { fileURLToPath } from 'url';
const __dirname = dirname(fileURLToPath(import.meta.url));
const outPath = resolve(__dirname, '_seo_posts_live.json');

function mapPosts(arr, corpus) {
  return (arr || [])
    .filter((d) => d && (d._type === 'post' || d.seoTitle || d.slug?.current))
    .filter((d) => d.title && (d.slug?.current || d.slug))
    .map((d) => ({
      id: d._id,
      title: d.title,
      slug: d.slug?.current || d.slug,
      seoTitle: d.seoTitle || '',
      seoDescription: d.seoDescription || '',
      corpus,
    }));
}
const seeded = mapPosts(posts, 'seeded (content+extra+finance-os)');
const pending = [
  ...mapPosts(july29Posts, 'working-tree untracked july29-posts.mjs'),
  ...mapPosts(july30Posts, 'working-tree untracked july30-posts.mjs'),
];
const glossary = (glossaryTerms || []).map((t) => ({
  id: t._id,
  term: t.term,
  slug: t.slug?.current || t.slug,
  shortDefinition: t.shortDefinition || '',
}));
writeFileSync(outPath, JSON.stringify({ seeded, pending, glossary }, null, 2));
console.log('seeded', seeded.length, 'pending', pending.length, 'glossary', glossary.length);
"""
    node_path = ROOT / "scripts" / "_seo_node_dump.mjs"
    node_path.write_text(node_script, encoding="utf-8")
    subprocess.run(["node", str(node_path)], cwd=ROOT / "scripts", check=True)
    live = json.loads(
        (ROOT / "scripts" / "_seo_posts_live.json").read_text(encoding="utf-8")
    )
    after_posts = live["seeded"]
    pending_posts = live["pending"]
    after_glossary = live.get("glossary") or []
    if len(after_glossary) < 15:
        current_seed_text = (ROOT / "frontend/sanity/seed/content.mjs").read_text(
            encoding="utf-8"
        )
        after_glossary = extract_glossary(current_seed_text)

    # --- Build workbook ---
    wb = Workbook()

    # Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "SMPL.ai SEO Keywords — Before vs After (late July / early Aug 2026)"
    ws["A1"].font = Font(bold=True, size=14)
    summary_lines = [
        "",
        "IMPORTANT: There was no formal pre-change keyword target list in the repo.",
        "The BEFORE column is RECONSTRUCTED from git history (commits around 2026-07-21),",
        "Sanity seed files, and marketing page metadata — not from a GSC/Ahrefs export.",
        "",
        "Date range of the shift:",
        "  • 2026-07-21 — 2655a8c: Align marketing SEO metadata with 'Rev 1.0 tag comparison'",
        "      (homepage title/description pivot: AI OS → SaaS FP&A Software & Board Reporting;",
        "       pipe | house style; unique titles/canonicals; login + board-sample noindex).",
        "  • 2026-07-21 — 8e5517f: Sanity blog + glossary Resources launched with 3 cornerstone posts",
        "      that used Lock/Freeze close language + Freeze pack glossary term.",
        "  • 2026-07-21 — cd012b8: Replaced those 3 posts with ARR waterfall / board reporting / AI commentary set.",
        "  • 2026-07-21 — 3577275: Removed freeze/lock IP jargon from glossary (Freeze pack, Combined scenario →",
        "      FP&A, Rolling forecast); glossary meta dropped 'freeze pack'.",
        "  • 2026-07-22 — 763cf52: Added Trust + SaaS close posts (CFO trust; spreadsheet reconciliation).",
        "  • 2026-07-27 — f0d6f48: GSC duplicate-without-canonical fix (per-URL canonicals).",
        "  • 2026-07-28 — 09834b6: Finance OS cluster posts (AI OS for SaaS finance; Finance OS vs FP&A).",
        "  • Working tree (untracked as of extract): july29-posts.mjs + july30-posts.mjs (4 more posts).",
        "",
        "Strategic intent shift (inferred from titles/meta, not a keyword brief):",
        "  BEFORE: Brand/category = 'AI operating system for SaaS finance' / 'CFO Operating Intelligence';",
        "          close IP language (Load→Validate→Lock→Freeze, freeze pack).",
        "  AFTER:  Ranking language = 'SaaS FP&A software', board reporting, trust/traceability, close;",
        "          Finance OS cluster retained as thought-leadership; Lock/Freeze public jargon retired.",
        "",
        "Sources used: git show on site.ts, marketing layouts, sanity/seed/*; docs/SANITY_BLOG_SETUP.md;",
        "agent transcript 8727df89 (blog/glossary implement) — no separate keyword matrix found there.",
        "Not SOC 2 certified. File is for Matt's internal marketing comparison only.",
        "",
        f"Output path: {OUT}",
    ]
    for i, line in enumerate(summary_lines, start=2):
        ws[f"A{i}"] = line
    ws.column_dimensions["A"].width = 110

    # Before sheet
    ws_b = wb.create_sheet("Before")
    before_rows = []
    for p in before_pages:
        before_rows.append(
            [
                "Page meta",
                p["surface"],
                p["url_path"],
                "",
                p["title"],
                p["meta_description"],
                p["primary_keyword"],
                p["source"],
            ]
        )
    before_rows.append(
        [
            "Page meta",
            before_glossary_meta["surface"],
            before_glossary_meta["url_path"],
            "",
            before_glossary_meta["title"],
            before_glossary_meta["meta_description"],
            before_glossary_meta["primary_keyword"],
            before_glossary_meta["source"],
        ]
    )
    for p in before_posts:
        before_rows.append(
            [
                "Blog post",
                p["title"],
                f"/blog/{p['slug']}",
                p["slug"],
                p.get("seoTitle") or p["title"],
                p.get("seoDescription") or "",
                infer_primary_keyword(p["title"], p.get("seoTitle", "")),
                "frontend/sanity/seed/content.mjs @ 8e5517f (retired in cd012b8 / seed-sanity.mjs)",
            ]
        )
    for t in before_glossary:
        before_rows.append(
            [
                "Glossary term",
                t["term"],
                f"/glossary/{t['slug']}",
                t["slug"],
                f"{t['term']} | SMPL.ai (term page)",
                t.get("shortDefinition") or "",
                t["term"],
                "glossary seed @ 8e5517f (freeze-pack / combined-scenario later removed)",
            ]
        )
    write_rows(
        ws_b,
        [
            "Type",
            "Label / H1",
            "URL path",
            "Slug",
            "Title / SEO title",
            "Meta / short definition",
            "Primary keyword (inferred)",
            "Source (file + commit)",
        ],
        before_rows,
    )

    # After sheet
    ws_a = wb.create_sheet("After")
    after_rows = []
    for p in after_pages:
        after_rows.append(
            [
                "Page meta",
                p["surface"],
                p["url_path"],
                "",
                p["title"],
                p["meta_description"],
                p["primary_keyword"],
                p["source"],
            ]
        )
    for p in after_posts:
        after_rows.append(
            [
                "Blog post",
                p["title"],
                f"/blog/{p['slug']}",
                p["slug"],
                p.get("seoTitle") or p["title"],
                p.get("seoDescription") or "",
                infer_primary_keyword(p["title"], p.get("seoTitle", "")),
                p.get("corpus") or "frontend/sanity/seed (current)",
            ]
        )
    for p in pending_posts:
        after_rows.append(
            [
                "Blog post (untracked seed)",
                p["title"],
                f"/blog/{p['slug']}",
                p["slug"],
                p.get("seoTitle") or p["title"],
                p.get("seoDescription") or "",
                infer_primary_keyword(p["title"], p.get("seoTitle", "")),
                p.get("corpus") or "untracked",
            ]
        )
    for t in after_glossary:
        after_rows.append(
            [
                "Glossary term",
                t["term"],
                f"/glossary/{t['slug']}",
                t["slug"],
                f"{t['term']} | SMPL.ai (term page)",
                t.get("shortDefinition") or "",
                t["term"],
                "frontend/sanity/seed/content.mjs (current)",
            ]
        )
    write_rows(
        ws_a,
        [
            "Type",
            "Label / H1",
            "URL path",
            "Slug",
            "Title / SEO title",
            "Meta / short definition",
            "Primary keyword (inferred)",
            "Source",
        ],
        after_rows,
    )

    # Diff sheet
    ws_d = wb.create_sheet("Diff")
    diff_rows = []

    # Page-level side by side
    page_pairs = [
        (
            "/",
            "SMPL.ai · AI operating system for SaaS finance teams",
            "SaaS FP&A Software & Board Reporting, Built to Be Trusted | SMPL.ai",
            "Changed",
            "AI operating system → SaaS FP&A software + board reporting / trust",
        ),
        (
            "/book-demo",
            "Book a demo · SMPL.ai",
            "Book a Demo | SMPL.ai SaaS FP&A Platform",
            "Changed",
            "Added SaaS FP&A platform keywording",
        ),
        (
            "/pricing",
            "(no unique meta)",
            "SMPL.ai Pricing | SaaS FP&A Plans",
            "Added",
            "First dedicated pricing SEO title/description",
        ),
        (
            "/request-quote",
            "Request a quote · SMPL.ai",
            "Request a Quote | SMPL.ai Pricing & Packaging",
            "Changed",
            "Added FP&A / packaging intent",
        ),
        (
            "/glossary",
            "... MD&A, freeze pack, and more.",
            "... MD&A, FP&A, and more.",
            "Changed",
            "Removed freeze pack from glossary meta; added FP&A",
        ),
        (
            "/blog",
            "(did not exist before 2026-07-21)",
            "Blog | SMPL.ai — SaaS FP&A, board reporting, close, AI commentary",
            "Added",
            "New Resources surface",
        ),
    ]
    for path, before, after, status, note in page_pairs:
        diff_rows.append(
            ["Page meta", path, before, after, status, note, path.strip("/").replace("/", "-") or "homepage"]
        )

    before_slugs = {p["slug"]: p for p in before_posts}
    after_slugs = {p["slug"]: p for p in after_posts}
    pending_slugs = {p["slug"]: p for p in pending_posts}
    all_after = {**after_slugs, **pending_slugs}

    for slug, p in before_slugs.items():
        if slug not in all_after:
            diff_rows.append(
                [
                    "Blog post",
                    f"/blog/{slug}",
                    p.get("seoTitle") or p["title"],
                    "(removed from seed; retiredPostIds in seed-sanity.mjs)",
                    "Removed",
                    "Lock/Freeze / early cornerstone cluster retired 2026-07-21 (cd012b8)",
                    infer_primary_keyword(p["title"], p.get("seoTitle", "")),
                ]
            )
        else:
            a = all_after[slug]
            diff_rows.append(
                [
                    "Blog post",
                    f"/blog/{slug}",
                    p.get("seoTitle") or p["title"],
                    a.get("seoTitle") or a["title"],
                    "Kept",
                    "",
                    infer_primary_keyword(a["title"], a.get("seoTitle", "")),
                ]
            )

    for slug, p in after_slugs.items():
        if slug not in before_slugs:
            diff_rows.append(
                [
                    "Blog post",
                    f"/blog/{slug}",
                    "(not in initial seed)",
                    p.get("seoTitle") or p["title"],
                    "Added",
                    "In current seed (content / extra / finance-os)",
                    infer_primary_keyword(p["title"], p.get("seoTitle", "")),
                ]
            )
    for slug, p in pending_slugs.items():
        if slug not in before_slugs and slug not in after_slugs:
            diff_rows.append(
                [
                    "Blog post",
                    f"/blog/{slug}",
                    "(not in initial seed)",
                    p.get("seoTitle") or p["title"],
                    "Added (untracked seed file)",
                    "july29/july30 working-tree files — confirm published to Sanity",
                    infer_primary_keyword(p["title"], p.get("seoTitle", "")),
                ]
            )

    before_g = {t["slug"]: t for t in before_glossary}
    after_g = {t["slug"]: t for t in after_glossary}
    for slug, t in before_g.items():
        if slug not in after_g:
            diff_rows.append(
                [
                    "Glossary term",
                    f"/glossary/{slug}",
                    t["term"],
                    "(removed)",
                    "Removed",
                    "3577275 removed freeze/lock IP jargon terms",
                    t["term"],
                ]
            )
        else:
            diff_rows.append(
                [
                    "Glossary term",
                    f"/glossary/{slug}",
                    t["term"],
                    after_g[slug]["term"],
                    "Kept",
                    "",
                    t["term"],
                ]
            )
    for slug, t in after_g.items():
        if slug not in before_g:
            diff_rows.append(
                [
                    "Glossary term",
                    f"/glossary/{slug}",
                    "(not present)",
                    t["term"],
                    "Added",
                    "3577275 added FP&A / Rolling forecast replacements",
                    t["term"],
                ]
            )

    # Keyword theme rollup
    theme_changes = [
        (
            "Theme / keyword",
            "AI operating system for SaaS finance",
            "Still used in Finance OS cluster + landing hero copy, but NOT homepage <title>",
            "Reprioritized",
            "Homepage title pivoted to SaaS FP&A Software & Board Reporting",
            "AI operating system for SaaS finance",
        ),
        (
            "Theme / keyword",
            "CFO Operating Intelligence",
            "(dropped from default title)",
            "Removed",
            "Was root layout default title before 2655a8c",
            "CFO Operating Intelligence",
        ),
        (
            "Theme / keyword",
            "Load → Validate → Lock → Freeze / freeze pack",
            "Internal product language only; not public SEO glossary targets",
            "Removed (public)",
            "Posts + glossary terms retired; close content reframed",
            "SaaS close Lock Freeze",
        ),
        (
            "Theme / keyword",
            "(weak brand-only page titles)",
            "SaaS FP&A software / plans / platform in titles",
            "Added",
            "Rev 1.0 tag comparison alignment",
            "SaaS FP&A software",
        ),
        (
            "Theme / keyword",
            "(no Finance OS blog cluster)",
            "AI Operating System for SaaS Finance; Finance OS vs FP&A Software",
            "Added",
            "finance-os-posts.mjs @ 09834b6 2026-07-28",
            "Finance OS / AI OS for SaaS finance",
        ),
        (
            "Theme / keyword",
            "(no ARR waterfall / board trust cluster)",
            "ARR Waterfall vs GAAP; board reporting breakdown; AI variance; CFO trust; spreadsheet recon",
            "Added",
            "cd012b8 + 763cf52 marketing articles",
            "ARR waterfall; SaaS board reporting; AI variance commentary",
        ),
    ]
    for row in theme_changes:
        diff_rows.append(
            [row[0], "(sitewide)", row[1], row[2], row[3], row[4], row[5]]
        )

    write_rows(
        ws_d,
        [
            "Type",
            "URL / scope",
            "Before",
            "After",
            "Status",
            "Notes",
            "Primary keyword / intent",
        ],
        diff_rows,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT}")
    print(f"Before posts: {len(before_posts)}; After seeded: {len(after_posts)}; Pending: {len(pending_posts)}")
    print(f"Before glossary: {len(before_glossary)}; After glossary: {len(after_glossary)}")

    # cleanup temp helpers
    for p in [
        ROOT / "scripts" / "_seo_node_dump.mjs",
        ROOT / "scripts" / "_seo_posts_live.json",
        ROOT / "scripts" / "_seo_extract.json",
        ROOT / "scripts" / "_extract_seo_tmp.py",
    ]:
        if p.exists():
            p.unlink()


if __name__ == "__main__":
    main()
